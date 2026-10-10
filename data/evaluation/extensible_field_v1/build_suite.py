#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""设备维保工单领域（未参与开发的领域）的声明式配置与材料生成器（WP5a / B2）。

本文件是**冻结件**：`manifest.json`、`labels.json`、`materials/*` 都由它生成，
且生成过程只依赖下面手写的 `FIELDS` / `MATERIALS`（作者标注的期望值）。

纪律（与仓库既有冻结评测一致）
------------------------------
1. **先写期望再评测**：材料内容与 `labels.json` 的期望值都来自本文件的 SPEC，
   与任何抽取器输出无关；抽取器（`scripts/eval_extensible_fields.py`）只读这些冻结件。
2. **确定性**：`.docx` / `.xlsx` 本质是 zip，python-docx / openpyxl 默认写入
   “当前时间”的 zip 时间戳与 core properties，逐字节不可复现。这里把 zip 条目
   时间戳与文档 core 时间统一归一化到 2026-01-01，因此同一 SPEC 重复生成得到
   **逐字节相同**的文件（由 tests/test_extensible_fields.py 断言）。
3. **只读校验**：`--check` 重新计算材料字节与标签，与磁盘上的冻结件比对；
   材料被篡改、标签被篡改、摘要被篡改都会以退出码 1 失败。
4. **不写库**：本脚本不连接任何数据库，不新增迁移，不改元数据库表结构。

用法
----
    .venv/bin/python data/evaluation/extensible_field_v1/build_suite.py --write
    .venv/bin/python data/evaluation/extensible_field_v1/build_suite.py --check
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[3]
SUITE_DIR = Path(__file__).resolve().parent

FIXED_DATE = (2026, 1, 1, 0, 0, 0)
FIXED_CORE_DATE = "2026-01-01T00:00:00Z"

SUITE_VERSION = "extensible-field-v1"
DATA_VERSION = "maintenance-v1"
DOMAIN = "设备维保工单"

#: 声明式字段定义。属性名刻意与元数据库注册表 `schema_columns` 对齐
#: （business_name / aliases / enum_map / semantic_role / sensitivity / unit），
#: 这样新领域的配置行与既有注册表行的属性契约一致，无需改表结构。
FIELDS = [
    {"name": "equipment_id", "label": "设备编号", "aliases": ["设备号", "设备ID"],
     "value_type": "string", "pattern": r"^EQ-\d{3}$", "enum_map": {},
     "semantic_role": "id", "sensitivity": "public", "unit": None, "required": True},
    {"name": "fault_code", "label": "故障代码", "aliases": ["故障码", "故障编码"],
     "value_type": "string", "pattern": r"^F-\d{2}$", "enum_map": {},
     "semantic_role": "dimension", "sensitivity": "public", "unit": None, "required": True},
    {"name": "service_date", "label": "维保日期", "aliases": ["服务日期", "维保时间"],
     "value_type": "date", "pattern": r"^\d{4}-\d{2}-\d{2}$", "enum_map": {},
     "semantic_role": "time", "sensitivity": "public", "unit": None, "required": True},
    {"name": "labor_hours", "label": "工时", "aliases": ["维修工时", "人工工时"],
     "value_type": "number", "pattern": r"^\d+(\.\d+)?$", "enum_map": {},
     "semantic_role": "metric", "sensitivity": "public", "unit": "小时", "required": True},
    {"name": "spare_part_no", "label": "备件号", "aliases": ["备件编号", "更换备件"],
     "value_type": "string", "pattern": r"^SP-\d{3}$", "enum_map": {},
     "semantic_role": "dimension", "sensitivity": "public", "unit": None, "required": False},
    {"name": "conclusion", "label": "维保结论", "aliases": ["处理结论", "维保结果"],
     "value_type": "enum", "pattern": r"^(正常|需更换备件|需复检|观察)$",
     "enum_map": {"normal": "正常", "replace_part": "需更换备件",
                  "reinspect": "需复检", "monitor": "观察"},
     "semantic_role": "dimension", "sensitivity": "public", "unit": None, "required": True},
]

FIELD_ORDER = [f["name"] for f in FIELDS]

#: 材料与作者标注的期望值。字段缺省（键不存在）表示“材料里没有这个字段”，
#: 期望值为 null —— 抽取器必须如实返回“无”，不得从别处补一个值。
MATERIALS = [
    {
        "material_id": "m01_txt_labeled",
        "file": "materials/m01_txt_labeled.txt",
        "format": "txt", "mode": "text_blocks",
        "record_separator": "", "kv_separator": "：",
        "labels": {"equipment_id": "设备编号", "fault_code": "故障代码", "service_date": "维保日期",
                   "labor_hours": "工时", "spare_part_no": "备件号", "conclusion": "维保结论"},
        "records": [
            {"equipment_id": "EQ-001", "fault_code": "F-01", "service_date": "2026-03-02",
             "labor_hours": "2.5", "spare_part_no": "SP-011", "conclusion": "需更换备件"},
            {"equipment_id": "EQ-002", "fault_code": "F-04", "service_date": "2026-03-03",
             "labor_hours": "1.0", "spare_part_no": "SP-012", "conclusion": "正常"},
            {"equipment_id": "EQ-007", "fault_code": "F-12", "service_date": "2026-03-05",
             "labor_hours": "4.5", "spare_part_no": "SP-013", "conclusion": "需复检"},
        ],
    },
    {
        "material_id": "m02_docx_labeled",
        "file": "materials/m02_docx_labeled.docx",
        "format": "docx", "mode": "text_blocks",
        "record_separator": "----", "kv_separator": "：",
        "labels": {"equipment_id": "设备号", "fault_code": "故障码", "service_date": "服务日期",
                   "labor_hours": "维修工时", "spare_part_no": "备件编号", "conclusion": "处理结论"},
        "records": [
            {"equipment_id": "EQ-003", "fault_code": "F-02", "service_date": "2026-03-06",
             "labor_hours": "3.0", "spare_part_no": "SP-014", "conclusion": "需更换备件"},
            {"equipment_id": "EQ-004", "fault_code": "F-07", "service_date": "2026-03-07",
             "labor_hours": "0.5", "conclusion": "正常"},
            {"equipment_id": "EQ-009", "fault_code": "F-03", "service_date": "2026-03-09",
             "labor_hours": "6.0", "spare_part_no": "SP-015", "conclusion": "观察"},
        ],
    },
    {
        "material_id": "m03_xlsx_table",
        "file": "materials/m03_xlsx_table.xlsx",
        "format": "xlsx", "mode": "xlsx_table",
        "sheet": "维保工单", "header_row": 1,
        "labels": {"equipment_id": "设备编号", "fault_code": "故障代码", "service_date": "维保日期",
                   "labor_hours": "工时", "spare_part_no": "备件号", "conclusion": "维保结论"},
        "records": [
            {"equipment_id": "EQ-005", "fault_code": "F-05", "service_date": "2026-03-10",
             "labor_hours": "2.0", "spare_part_no": "SP-016", "conclusion": "正常"},
            {"equipment_id": "EQ-006", "fault_code": "F-06", "service_date": "2026-03-11",
             "labor_hours": "3.5", "spare_part_no": "SP-017", "conclusion": "需更换备件"},
            {"equipment_id": "EQ-010", "fault_code": "F-08", "service_date": "2026-03-12",
             "labor_hours": "1.5", "spare_part_no": "SP-018", "conclusion": "需复检"},
            {"equipment_id": "EQ-011", "fault_code": "F-09", "service_date": "2026-03-13",
             "labor_hours": "5.0", "spare_part_no": "SP-019", "conclusion": "观察"},
        ],
    },
    {
        "material_id": "m04_docx_table",
        "file": "materials/m04_docx_table.docx",
        "format": "docx", "mode": "docx_table",
        "labels": {"equipment_id": "设备ID", "fault_code": "故障编码", "service_date": "维保时间",
                   "labor_hours": "人工工时", "spare_part_no": "更换备件", "conclusion": "维保结果"},
        "records": [
            {"equipment_id": "EQ-012", "fault_code": "F-10", "service_date": "2026-03-14",
             "labor_hours": "2.25", "spare_part_no": "SP-020", "conclusion": "正常"},
            {"equipment_id": "EQ-013", "fault_code": "F-02", "service_date": "2026-03-15",
             "labor_hours": "0.75", "spare_part_no": "SP-021", "conclusion": "需更换备件"},
            {"equipment_id": "EQ-014", "fault_code": "F-11", "service_date": "2026-03-16",
             "labor_hours": "3.75", "conclusion": "正常"},
        ],
    },
    {
        "material_id": "m05_txt_delimited",
        "file": "materials/m05_txt_delimited.txt",
        "format": "txt", "mode": "text_blocks",
        "record_separator": "", "kv_separator": "|",
        "labels": {"equipment_id": "设备编号", "fault_code": "故障代码", "service_date": "维保日期",
                   "labor_hours": "工时", "spare_part_no": "备件号", "conclusion": "维保结论"},
        "records": [
            {"equipment_id": "EQ-015", "fault_code": "F-12", "service_date": "2026-03-17",
             "labor_hours": "1.0", "spare_part_no": "SP-022", "conclusion": "需复检"},
            {"equipment_id": "EQ-016", "fault_code": "F-13", "service_date": "2026-03-18",
             "labor_hours": "4.0", "spare_part_no": "SP-023", "conclusion": "观察"},
            {"equipment_id": "EQ-017", "fault_code": "F-14", "service_date": "2026-03-19",
             "labor_hours": "2.0", "spare_part_no": "SP-024", "conclusion": "需更换备件"},
        ],
    },
    {
        "material_id": "m06_xlsx_two_sheets",
        "file": "materials/m06_xlsx_two_sheets.xlsx",
        "format": "xlsx", "mode": "xlsx_table",
        "sheets": ["维保工单一", "维保工单二"], "header_row": 1,
        "labels": {"equipment_id": "设备号", "fault_code": "故障代码", "service_date": "服务日期",
                   "labor_hours": "工时", "spare_part_no": "备件编号", "conclusion": "处理结论"},
        "records": [
            {"equipment_id": "EQ-018", "fault_code": "F-15", "service_date": "2026-03-20",
             "labor_hours": "3.0", "spare_part_no": "SP-025", "conclusion": "正常"},
            {"equipment_id": "EQ-019", "fault_code": "F-16", "service_date": "2026-03-21",
             "spare_part_no": "SP-026", "conclusion": "需复检"},
            {"equipment_id": "EQ-020", "fault_code": "F-17", "service_date": "2026-03-22",
             "labor_hours": "1.25", "spare_part_no": "SP-027", "conclusion": "观察"},
            {"equipment_id": "EQ-021", "fault_code": "F-18", "service_date": "2026-03-23",
             "labor_hours": "5.5", "spare_part_no": "SP-028", "conclusion": "需更换备件"},
        ],
    },
]

#: 三档对照的“示例”定义：显式声明的 `material_id/record_key`，保证三档可复现。
EXAMPLE_TIERS = {
    "0_examples": [],
    "3_examples": ["m01_txt_labeled/EQ-001", "m01_txt_labeled/EQ-002", "m03_xlsx_table/EQ-005"],
    "10_examples": ["m01_txt_labeled/EQ-001", "m01_txt_labeled/EQ-002", "m01_txt_labeled/EQ-007",
                    "m02_docx_labeled/EQ-003", "m02_docx_labeled/EQ-004", "m03_xlsx_table/EQ-005",
                    "m03_xlsx_table/EQ-006", "m04_docx_table/EQ-012", "m05_txt_delimited/EQ-015",
                    "m06_xlsx_two_sheets/EQ-018"],
}


# ---------------------------------------------------------------- 生成（纯函数）

def _normalize_zip(data: bytes) -> bytes:
    """把 zip 条目时间戳/顺序/权限与 core.xml 的 modified 时间归一化。

    python-docx / openpyxl 会写入“当前时间”：zip 条目时间戳来自 `time.localtime()`，
    而 openpyxl 还会在 `write_workbook()` 里强行覆盖 `properties.modified`
    （`.venv/.../openpyxl/writer/excel.py:292`）。因此仅设置 properties 不足以复现；
    这里统一改写 docProps/core.xml 的 `dcterms:modified`，使同一 SPEC 逐字节可复现。
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = _normalize_core_properties(payload)
            info = zipfile.ZipInfo(name, date_time=FIXED_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, payload)
    return out.getvalue()


def _normalize_core_properties(payload: bytes) -> bytes:
    import re
    text = payload.decode("utf-8")
    for tag in ("created", "modified"):
        text = re.sub(rf"(<dcterms:{tag}[^>]*>)[^<]*(</dcterms:{tag}>)",
                      rf"\g<1>{FIXED_CORE_DATE}\g<2>", text)
    return text.encode("utf-8")


def _record_key(material: dict, record: dict, index: int) -> str:
    return record.get("equipment_id") or f"{material['material_id']}#{index + 1}"


def material_records(material: dict) -> list[dict]:
    """材料里实际出现的字段（缺省字段不补值）。"""
    return [{key: value for key, value in record.items() if value is not None}
            for record in material["records"]]


def text_lines(material: dict, record: dict) -> list[str]:
    separator = material["kv_separator"]
    return [f"{material['labels'][field]}{separator}{record[field]}"
            for field in FIELD_ORDER if field in record and record[field] is not None]


def material_bytes(material: dict) -> bytes:
    """按 `mode` 声明确定性生成材料字节。"""
    if material["format"] == "txt":
        chunks = []
        for record in material["records"]:
            chunks.append("\n".join(text_lines(material, record)))
        joiner = "\n\n" if material["record_separator"] == "" else f"\n{material['record_separator']}\n"
        return (joiner.join(chunks) + "\n").encode("utf-8")

    if material["format"] == "docx":
        from docx import Document
        document = Document()
        document.core_properties.created = _fixed_datetime()
        document.core_properties.modified = _fixed_datetime()
        if material["mode"] == "docx_table":
            table = document.add_table(rows=1, cols=len(FIELD_ORDER))
            table.style = "Table Grid"
            for index, field in enumerate(FIELD_ORDER):
                table.rows[0].cells[index].text = material["labels"][field]
            for record in material["records"]:
                cells = table.add_row().cells
                for index, field in enumerate(FIELD_ORDER):
                    value = record.get(field)
                    cells[index].text = "" if value is None else str(value)
        else:
            blocks = []
            for record in material["records"]:
                blocks.append(text_lines(material, record))
            body = []
            for block_index, block in enumerate(blocks):
                if block_index:
                    body.append(material["record_separator"])
                body.extend(block)
            for line in body:
                document.add_paragraph(line)
        buffer = io.BytesIO()
        document.save(buffer)
        return _normalize_zip(buffer.getvalue())

    if material["format"] == "xlsx":
        import openpyxl
        workbook = openpyxl.Workbook()
        workbook.properties.created = _fixed_datetime()
        workbook.properties.modified = _fixed_datetime()
        sheets = material.get("sheets") or [material["sheet"]]
        records = material["records"]
        # 两个 sheet 时按声明顺序均分记录（第一个 sheet 取下整，其余递增）。
        if len(sheets) == 1:
            groups = [records]
        else:
            size = len(records) // len(sheets)
            groups = [records[index * size:(index + 1) * size] for index in range(len(sheets) - 1)]
            groups.append(records[(len(sheets) - 1) * size:])
        for index, sheet_name in enumerate(sheets):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = sheet_name
            sheet.append([material["labels"][field] for field in FIELD_ORDER])
            for record in groups[index]:
                sheet.append([record.get(field) for field in FIELD_ORDER])
        buffer = io.BytesIO()
        workbook.save(buffer)
        return _normalize_zip(buffer.getvalue())

    raise ValueError(f"unsupported format: {material['format']}")


def _fixed_datetime():
    import datetime
    return datetime.datetime(2026, 1, 1)


def labels_payload() -> dict:
    assertions = []
    for material in MATERIALS:
        for index, record in enumerate(material["records"]):
            key = _record_key(material, record, index)
            for field in FIELD_ORDER:
                value = record.get(field)
                assertions.append({
                    "assertion_id": f"{material['material_id']}/{key}/{field}",
                    "material_id": material["material_id"],
                    "record_key": key,
                    "field": field,
                    "expected": None if value is None else str(value),
                })
    return {
        "suite_version": SUITE_VERSION,
        "data_version": DATA_VERSION,
        "domain": DOMAIN,
        "frozen_before_evaluation": True,
        "annotation": "期望值来自 build_suite.py 中作者标注的 MATERIALS（材料 SPEC），"
                      "与 scripts/eval_extensible_fields.py 的输出无关；缺省字段期望为 null（不得补值）。",
        "assertions": assertions,
    }


def labels_bytes() -> bytes:
    return (json.dumps(labels_payload(), ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def manifest_payload() -> dict:
    materials = []
    for material in MATERIALS:
        assert material["material_id"] not in [m["material_id"] for m in materials]
        materials.append({
            "material_id": material["material_id"],
            "file": material["file"],
            "format": material["format"],
            "mode": material["mode"],
            "record_separator": material.get("record_separator"),
            "kv_separator": material.get("kv_separator"),
            "sheet": material.get("sheet"),
            "sheets": material.get("sheets"),
            "header_row": material.get("header_row"),
            "labels": material["labels"],
            "record_count": len(material["records"]),
            "labeled_record_count": len(material["records"]),
            "sha256": hashlib.sha256(material_bytes(material)).hexdigest(),
        })
    return {
        "suite_version": SUITE_VERSION,
        "data_version": DATA_VERSION,
        "domain": DOMAIN,
        "frozen": True,
        "registry_contract": {
            "reuses": "schema_columns 属性契约（business_name/aliases/enum_map/semantic_role/sensitivity/unit）",
            "ddl_changes": 0,
            "migrations_added": 0,
            "registry_writes": 0,
        },
        "record_key_field": "equipment_id",
        "field_count": len(FIELDS),
        "fields": FIELDS,
        "materials": materials,
        "example_tiers": EXAMPLE_TIERS,
        "labels_file": "labels.json",
        "labels_sha256": hashlib.sha256(labels_bytes()).hexdigest(),
        "assertion_count": len(labels_payload()["assertions"]),
        "generator": "data/evaluation/extensible_field_v1/build_suite.py",
        "normalization": "zip 时间戳与 core 时间统一为 2026-01-01；条目名排序、压缩方式固定",
    }


def manifest_bytes() -> bytes:
    return (json.dumps(manifest_payload(), ensure_ascii=False, indent=2) + "\n").encode("utf-8")


# ---------------------------------------------------------------- 校验与写入

def check(suite_dir: Path) -> list[str]:
    problems = []
    manifest_path = suite_dir / "manifest.json"
    labels_path = suite_dir / "labels.json"
    for path in (manifest_path, labels_path):
        if not path.is_file():
            problems.append(f"MISSING: {path}")
    if problems:
        return problems

    expected_manifest = manifest_bytes()
    if manifest_path.read_bytes() != expected_manifest:
        problems.append("MANIFEST_DRIFT: manifest.json 与 build_suite.py 的声明不一致")
    if labels_path.read_bytes() != labels_bytes():
        problems.append("LABELS_DRIFT: labels.json 与 SPEC 重算的期望值不一致")

    manifest = json.loads(manifest_path.read_bytes())
    for material in MATERIALS:
        path = suite_dir / material["file"]
        if not path.is_file():
            problems.append(f"MATERIAL_MISSING: {material['file']}")
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        declared = hashlib.sha256(material_bytes(material)).hexdigest()
        frozen = next((m["sha256"] for m in manifest["materials"]
                       if m["material_id"] == material["material_id"]), None)
        if digest != declared:
            problems.append(f"MATERIAL_TAMPERED: {material['file']} 磁盘内容与 SPEC 生成结果不一致")
        if frozen != declared:
            problems.append(f"MATERIAL_DIGEST_DRIFT: {material['material_id']} manifest 摘要与生成结果不一致")
    return problems


def write(suite_dir: Path) -> None:
    (suite_dir / "materials").mkdir(parents=True, exist_ok=True)
    for material in MATERIALS:
        (suite_dir / material["file"]).write_bytes(material_bytes(material))
    (suite_dir / "labels.json").write_bytes(labels_bytes())
    (suite_dir / "manifest.json").write_bytes(manifest_bytes())


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--write", action="store_true", help="按 SPEC 重新生成材料与冻结标签")
    parser.add_argument("--check", action="store_true", help="只读校验冻结件是否与 SPEC 一致")
    parser.add_argument("--suite-dir", type=Path, default=SUITE_DIR)
    args = parser.parse_args()
    if args.write == args.check:
        parser.error("必须且只能指定 --write 或 --check 之一")
    if args.write:
        write(args.suite_dir)
        print(f"wrote {len(MATERIALS)} materials + manifest.json + labels.json -> {args.suite_dir}")
        return 0
    problems = check(args.suite_dir)
    if problems:
        for problem in problems:
            print(f"FAIL {problem}")
        print(f"check failed: {len(problems)} problem(s)")
        return 1
    manifest = json.loads((args.suite_dir / "manifest.json").read_bytes())
    print(f"check ok: {len(MATERIALS)} materials, {manifest['field_count']} fields, "
          f"{manifest['assertion_count']} assertions, domain={manifest['domain']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
