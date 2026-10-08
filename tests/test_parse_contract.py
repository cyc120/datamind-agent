"""Unified parse/locator/relation contract: validation, adapters and PDF block mapping.

These tests are the review of the contract: they pin the documented constants, the
per-kind locator requirements, relation integrity, the ingest state machine, the adapter
registry (including explicit refusal) and the mapping of real ingest block shapes.
"""
import json
import re
import tempfile
import uuid
from pathlib import Path
import unittest
from unittest.mock import patch

from fastapi import HTTPException

from ingestion import ingest
from datamind import parse_contract as contract

ROOT = Path(__file__).resolve().parents[1]


def contract_block_columns():
    """The `document_contract_blocks` columns the checked-in migrations declare.

    Read from the migration files rather than the live database so the closure of the block
    shape can be checked without a database and before `db/migrate.py` has run.
    """
    columns = set()
    for path in sorted((ROOT / 'db/migrations').glob('*.sql')):
        text = path.read_text(encoding='utf-8')
        for body in re.findall(r'CREATE TABLE IF NOT EXISTS document_contract_blocks\s*\((.*?)\);',
                               text, re.S):
            for line in body.splitlines():
                token = line.strip().split(None, 1)
                if token and re.fullmatch(r'[a-z_]+', token[0]):
                    columns.add(token[0])
        columns.update(re.findall(
            r'ALTER TABLE document_contract_blocks\s+ADD COLUMN IF NOT EXISTS ([a-z_]+)', text))
    return columns


def live_contract_block_columns():
    """The columns the running database actually has (empty when it is unreachable)."""
    import psycopg2
    from datamind.agent import _write_dsn
    with psycopg2.connect(_write_dsn()) as conn, conn.cursor() as cur:
        cur.execute("""SELECT column_name FROM information_schema.columns
            WHERE table_name='document_contract_blocks'""")
        return {row[0] for row in cur.fetchall()}


def locator(kind='pdf_region', **overrides):
    base = {'locator_kind': kind, 'document_id': 'doc-1', 'source_version': 'v1'}
    if kind == 'pdf_region':
        base.update({'page_no': 1, 'page_width': 595.0, 'page_height': 842.0,
                     'bbox': [10.0, 20.0, 100.0, 120.0], 'unit': 'point',
                     'coordinate_system': 'top-left', 'rotation_degrees': 0})
    elif kind == 'image_region':
        base.update({'image_id': 'img-1', 'image_width': 800, 'image_height': 600,
                     'bbox': [0.0, 0.0, 100.0, 50.0], 'unit': 'px', 'coordinate_system': 'top-left'})
    elif kind == 'office_object':
        base.update({'part': 'word/document.xml', 'paragraph': 3})
    elif kind == 'sheet_range':
        base.update({'sheet': 'Sheet1', 'cell': 'B2'})
    elif kind == 'cad_entity':
        base.update({'layout': 'Model', 'handle': '2AF', 'layer': '0', 'model_xy': [0.0, 0.0], 'unit': 'mm'})
    elif kind == 'temporal_segment':
        base.update({'start': 0.0, 'end': 1.5, 'time_unit': 's', 'time_base': 'media'})
    elif kind == 'model_spatial_object':
        base.update({'object_id': 'obj-1', 'coordinate_reference': 'site', 'unit': 'm'})
    base.update(overrides)
    return base


def block(block_id='b1', locator_payload=None, **overrides):
    payload = {'block_id': block_id, 'document_id': 'doc-1', 'source_version': 'v1',
               'parser_version': 'pdf-ocr-0.4', 'block_type': 'paragraph', 'text': '原文',
               'native_attributes': {}, 'reading_order': 0,
               'locator': locator_payload or locator(), 'quality': {},
               'content_hash': 'a' * 64}
    payload.update(overrides)
    return payload


def relation(source='b1', target='b2', **overrides):
    payload = {'source_block_id': source, 'target_block_id': target,
               'relation_type': 'contains', 'evidence': '原文依据',
               'recovery_method': 'explicit_marker', 'state': 'candidate'}
    payload.update(overrides)
    return payload


def parse_result(**overrides):
    payload = {'classification': {'detected_media_type': 'application/pdf'},
               'adapter_capabilities': ['native_text'],
               'blocks': [block('b1'), block('b2')],
               'relations': [relation()],
               'parsed_scope': [{'kind': 'page', 'ref': 1, 'status': 'parsed'}],
               'failed_scope': [],
               'quality_findings': [],
               'unsupported_features': [],
               'artifact_refs': [],
               'manifest': {'parser_version': 'pdf-ocr-0.4'},
               'parse_status': 'parsed'}
    payload.update(overrides)
    return payload


class ContractConstantTests(unittest.TestCase):
    def test_documented_sets_are_pinned(self):
        self.assertEqual(contract.PARSE_STATUSES,
                         ('parsed', 'partial', 'needs_review', 'failed', 'unsupported'))
        self.assertEqual(contract.INGEST_STATES,
                         ('queued', 'classifying', 'parsing', 'validating', 'indexing', 'ready', 'failed'))
        self.assertEqual(contract.RELATION_TYPES,
                         ('precedes', 'contains', 'caption_of', 'header_of', 'continuation_of',
                          'annotates', 'connected_to', 'flows_to', 'references'))
        self.assertEqual(contract.RELATION_STATES, ('candidate', 'verified', 'needs_review'))
        self.assertEqual(contract.LOCATOR_KINDS,
                         ('pdf_region', 'image_region', 'office_object', 'sheet_range',
                          'cad_entity', 'temporal_segment', 'model_spatial_object'))

    def test_design_doc_path_is_relative_and_exists(self):
        path = contract.design_doc_path()
        self.assertTrue(path.is_file(), path)
        self.assertTrue(str(path).startswith(str(ROOT)))
        self.assertFalse(str(path).startswith('/Users/caicai/.codex'))


class IngestStateTests(unittest.TestCase):
    def test_every_documented_state_is_accepted(self):
        for state in contract.INGEST_STATES:
            self.assertEqual(contract.validate_ingest_state(state), state)

    def test_unknown_state_is_refused(self):
        for state in ('', 'READY', 'done', 3, None):
            with self.assertRaises(HTTPException) as caught:
                contract.validate_ingest_state(state)
            self.assertEqual(caught.exception.detail['error_code'], 'PARSE_STATUS_INVALID')

    def test_only_documented_transitions_are_allowed(self):
        self.assertEqual(contract.ingest_step('queued', 'classifying'), 'classifying')
        self.assertEqual(contract.ingest_step('indexing', 'ready'), 'ready')
        self.assertEqual(contract.ingest_step('parsing', 'failed'), 'failed')
        for state, target in (('queued', 'ready'), ('ready', 'parsing'), ('failed', 'parsing')):
            with self.assertRaises(HTTPException) as caught:
                contract.ingest_step(state, target)
            self.assertEqual(caught.exception.detail['error_code'], 'PARSE_STATUS_INVALID')


class LocatorTests(unittest.TestCase):
    def error(self, code, action):
        with self.assertRaises(HTTPException) as caught:
            action()
        self.assertEqual(caught.exception.detail['error_code'], code)

    def test_every_locator_kind_validates(self):
        for kind in contract.LOCATOR_KINDS:
            result = contract.validate_locator(locator(kind))
            self.assertEqual(result.locator_kind, kind)

    def test_pdf_region_bounds_and_page_number_are_enforced(self):
        self.error('PARSE_LOCATOR_INVALID',
                   lambda: contract.validate_locator(locator(page_no=0)))
        self.error('PARSE_LOCATOR_INVALID',
                   lambda: contract.validate_locator(locator(page_no=-1)))
        self.error('PARSE_LOCATOR_INVALID',
                   lambda: contract.validate_locator(locator(bbox=[100.0, 20.0, 10.0, 120.0])))
        self.error('PARSE_LOCATOR_INVALID',
                   lambda: contract.validate_locator(locator(bbox=[-5.0, 20.0, 100.0, 120.0])))
        self.error('PARSE_LOCATOR_INVALID',
                   lambda: contract.validate_locator(locator(bbox=[10.0, 20.0, 9999.0, 120.0])))

    def test_missing_required_fields_are_refused_per_kind(self):
        for kind in contract.LOCATOR_KINDS:
            for field in contract.LOCATOR_REQUIRED_FIELDS[kind]:
                payload = locator(kind)
                del payload[field]
                self.error('PARSE_LOCATOR_INVALID',
                           lambda payload=payload: contract.validate_locator(payload))

    def test_unknown_kind_is_refused(self):
        self.error('PARSE_LOCATOR_INVALID',
                   lambda: contract.validate_locator(locator(office_object=True, locator_kind='magic')))

    def test_cad_entity_handle_and_geometry_are_checked(self):
        self.error('PARSE_LOCATOR_INVALID',
                   lambda: contract.validate_locator(locator('cad_entity', handle='not-a-handle!')))
        self.assertIsNotNone(contract.validate_locator(locator('cad_entity', handle='2AF')))

    def test_pdf_region_helper_builds_a_valid_locator(self):
        built = contract.pdf_region_locator(document_id='doc-1', source_version='v1', page_no=2,
                                            bbox=[1.0, 2.0, 30.0, 40.0], page_width=595.0,
                                            page_height=842.0, rotation_degrees=90)
        self.assertEqual(built.locator_kind, 'pdf_region')
        self.assertEqual(built.payload['page_no'], 2)
        self.error('PARSE_LOCATOR_INVALID',
                   lambda: contract.pdf_region_locator(document_id='doc-1', source_version='v1',
                                                       page_no=0, bbox=[1.0, 2.0, 30.0, 40.0],
                                                       page_width=595.0, page_height=842.0,
                                                       rotation_degrees=0))


class RelationTests(unittest.TestCase):
    def error(self, code, action):
        with self.assertRaises(HTTPException) as caught:
            action()
        self.assertEqual(caught.exception.detail['error_code'], code)

    def test_relation_requires_known_endpoints_type_state_and_evidence(self):
        known = ['b1', 'b2']
        self.assertEqual(contract.validate_relation(relation(), known).relation_type, 'contains')
        for payload, code in ((relation(source='missing'), 'PARSE_RELATION_INVALID'),
                              (relation(target='missing'), 'PARSE_RELATION_INVALID'),
                              (relation(relation_type='magic'), 'PARSE_RELATION_INVALID'),
                              (relation(state='maybe'), 'PARSE_RELATION_INVALID')):
            self.error(code, lambda payload=payload: contract.validate_relation(payload, known))
        for field in contract.RELATION_REQUIRED_FIELDS:
            payload = relation()
            del payload[field]
            self.error('PARSE_RELATION_INVALID',
                       lambda payload=payload: contract.validate_relation(payload, known))

    def test_merge_by_similarity_is_not_an_acceptable_recovery_method(self):
        for method in contract.INSUFFICIENT_RECOVERY_METHODS:
            payload = relation(recovery_method=method)
            self.error('PARSE_RELATION_INVALID',
                       lambda payload=payload: contract.validate_relation(payload, ['b1', 'b2']))


class ParseResultTests(unittest.TestCase):
    def error(self, code, action):
        with self.assertRaises(HTTPException) as caught:
            action()
        self.assertEqual(caught.exception.detail['error_code'], code)

    def test_valid_result_round_trips(self):
        result = contract.validate_parse_result(parse_result())
        self.assertEqual(result.parse_status, 'parsed')
        self.assertEqual(len(result.blocks), 2)
        self.assertEqual(len(result.relations), 1)

    def test_every_required_field_is_enforced(self):
        for field in contract.RESULT_REQUIRED_FIELDS:
            payload = parse_result()
            del payload[field]
            self.error('PARSE_RESULT_INVALID', lambda payload=payload: contract.validate_parse_result(payload))

    def test_parse_status_must_be_documented(self):
        self.error('PARSE_RESULT_INVALID',
                   lambda: contract.validate_parse_result(parse_result(parse_status='done')))

    def test_relation_endpoints_must_exist_in_the_same_result(self):
        payload = parse_result(relations=[relation(source='b1', target='ghost')])
        self.error('PARSE_RELATION_INVALID', lambda: contract.validate_parse_result(payload))

    def test_unknown_extras_are_refused_rather_than_invented(self):
        payload = parse_result()
        payload['invented_field'] = 1
        self.error('PARSE_RESULT_INVALID', lambda: contract.validate_parse_result(payload))


class AdapterRegistryTests(unittest.TestCase):
    def setUp(self):
        contract._ADAPTERS.clear()

    def tearDown(self):
        contract._ADAPTERS.clear()

    def test_registration_lookup_and_capabilities(self):
        record = contract.register_adapter('pdf-native', media_types=['application/pdf'],
                                           capabilities=['native_text', 'tables'])
        self.assertEqual(record['name'], 'pdf-native')
        self.assertEqual(contract.adapter_for('application/pdf')['name'], 'pdf-native')
        self.assertIn('native_text', contract.capabilities('pdf-native'))
        self.assertIn('application/pdf', contract.supported_media_types())
        self.assertTrue(contract.supports_capability('pdf-native', 'tables'))
        self.assertEqual(contract.require_capability('pdf-native', 'tables'), 'tables')

    def test_unregistered_media_type_is_explicitly_refused(self):
        for media_type in ('image/tiff', '', 'nonsense'):
            with self.assertRaises(HTTPException) as caught:
                contract.adapter_for(media_type)
            self.assertEqual(caught.exception.detail['error_code'], 'PARSE_ADAPTER_UNKNOWN')

    def test_unsupported_capability_and_unknown_adapter_are_distinguished(self):
        contract.register_adapter('pdf-native', media_types=['application/pdf'],
                                  capabilities=['native_text'])
        with self.assertRaises(HTTPException) as caught:
            contract.require_capability('pdf-native', 'formula')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_ADAPTER_UNSUPPORTED')
        with self.assertRaises(HTTPException) as caught:
            contract.require_capability('missing-adapter', 'native_text')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_ADAPTER_UNKNOWN')

    def test_registration_inputs_are_validated(self):
        contract.register_adapter('pdf-native', media_types=['application/pdf'],
                                  capabilities=['native_text'])
        # The same media type cannot be claimed twice.
        with self.assertRaises(HTTPException) as caught:
            contract.register_adapter('other', media_types=['application/pdf'], capabilities=['x'])
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_RESULT_INVALID')
        for kwargs, code in (
                ({'media_types': [], 'capabilities': ['x']}, 'PARSE_RESULT_INVALID'),
                ({'media_types': ['application/pdf'], 'capabilities': []}, 'PARSE_RESULT_INVALID'),
                ({'media_types': ['application/pdf'], 'capabilities': ['']}, 'PARSE_RESULT_INVALID'),
                ({'media_types': ['nonsense'], 'capabilities': ['x']}, 'PARSE_ADAPTER_UNKNOWN')):
            with self.assertRaises(HTTPException) as caught:
                contract.register_adapter('bad-adapter', **kwargs)
            self.assertEqual(caught.exception.detail['error_code'], code)


class PdfMappingTests(unittest.TestCase):
    """Map the block dictionaries the ingest/OCR modules actually produce."""

    def native_block(self):
        return {'page_no': 1, 'bbox': [72.0, 100.0, 300.0, 140.0], 'block_type': 'paragraph',
                'heading_level': None, 'text': '正文内容', 'table_cells': None, 'ocr_confidence': None}

    def ocr_block(self):
        return {'page_no': 2, 'bbox': [10.0, 20.0, 200.0, 60.0], 'block_type': 'title',
                'heading_level': 2, 'text': '一、标题', 'table_cells': None, 'ocr_confidence': 0.92}

    def test_real_block_shapes_are_mapped_without_inventing_fields(self):
        blocks = [self.native_block(), self.ocr_block()]
        mapped = contract.from_pdf_blocks(blocks, page_sizes={1: (595.0, 842.0), 2: (595.0, 842.0)},
                                          parser_version='pdf-ocr-0.4', document_id='doc-1',
                                          source_version='v1')
        self.assertEqual(len(mapped), 2)
        first, second = mapped
        self.assertEqual(first['text'], '正文内容')
        self.assertEqual(first['page_no'], 1)
        self.assertIsNone(first['heading_level'])
        self.assertIsNone(first['ocr_confidence'])
        self.assertEqual(second['ocr_confidence'], 0.92)
        self.assertEqual(second['heading_level'], 2)
        for item in mapped:
            self.assertEqual(item['parser_version'], 'pdf-ocr-0.4')
            self.assertEqual(item['locator']['locator_kind'], 'pdf_region')
            self.assertEqual(item['locator']['bbox'], item['bbox'])
            self.assertEqual(item['locator']['page_width'], 595.0)
            self.assertIn('block_id', item)

    def test_missing_page_size_is_refused_instead_of_guessed(self):
        with self.assertRaises(HTTPException) as caught:
            contract.from_pdf_blocks([self.native_block()], page_sizes={}, parser_version='p',
                                     document_id='doc-1', source_version='v1')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_LOCATOR_INVALID')

    def test_mapped_blocks_are_accepted_by_the_result_contract(self):
        mapped = contract.from_pdf_blocks([self.native_block()], page_sizes={1: (595.0, 842.0)},
                                          parser_version='pdf-ocr-0.4', document_id='doc-1',
                                          source_version='v1')
        payload = parse_result(blocks=mapped, relations=[])
        self.assertEqual(len(contract.validate_parse_result(payload).blocks), 1)


if __name__ == '__main__':
    unittest.main()


class IngestIntegrationTests(unittest.TestCase):
    """The production parse path must actually emit a contract-validated result."""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.patch = patch.object(ingest, 'PAGE_IMG_DIR', Path(self.directory.name))
        self.patch.start()
        self.addCleanup(self.patch.stop)

    def native_pdf(self):
        return ROOT / 'data/knowledge_base/04_售后服务与退款管理制度.pdf'

    def scanned_pdf(self):
        return ROOT / 'data/knowledge_base/07_门店巡检记录_2025年12月_扫描版.pdf'

    def test_native_pdf_is_validated_against_the_contract(self):
        blocks, _, detail = ingest.parse_pdf(self.native_pdf(), 'contract-native', include_quality=True)
        summary = detail['parse_contract']
        self.assertEqual(summary['validation'], 'parsed', summary)
        self.assertEqual(summary['contract_version'], contract.CONTRACT_VERSION)
        self.assertEqual(summary['blocks'], len(blocks))
        self.assertEqual(summary['locator_kinds'], ['pdf_region'])
        self.assertEqual(summary['parse_status'], detail['parse_status'])
        # The result came through the registered adapter dispatch, not a direct call.
        self.assertEqual(summary['adapter'], 'pdf-native')

    def test_scanned_pdf_ocr_blocks_are_validated_too(self):
        blocks, _, detail = ingest.parse_pdf(self.scanned_pdf(), 'contract-scan', include_quality=True)
        summary = detail['parse_contract']
        self.assertEqual(summary['validation'], 'parsed', summary)
        self.assertEqual(summary['blocks'], len(blocks))
        self.assertTrue(blocks)

    def test_a_contract_violation_is_recorded_not_raised(self):
        import pymupdf
        path = self.native_pdf()
        with pymupdf.open(path) as document:
            page = document[0]
            oversized = {'page_no': 1, 'bbox': [0.0, 0.0, page.rect.width + 500.0, page.rect.height],
                         'block_type': 'paragraph', 'text': '越界块'}
            detail = {'source_sha256': 'a' * 64, 'parse_status': 'parsed', 'pages': []}
            summary = ingest._contract_summary(document, path, 'contract-bad', [oversized], detail)
        self.assertEqual(summary['validation'], 'failed', summary)
        self.assertEqual(summary['error_code'], 'PARSE_LOCATOR_INVALID')

    def test_no_blocks_skips_validation_instead_of_inventing_one(self):
        import pymupdf
        path = self.native_pdf()
        with pymupdf.open(path) as document:
            detail = {'source_sha256': 'b' * 64, 'parse_status': 'failed', 'pages': []}
            summary = ingest._contract_summary(document, path, 'contract-empty', [], detail)
        self.assertEqual(summary['validation'], 'skipped', summary)
        self.assertEqual(summary['reason'], 'no_blocks')


class RoutingTests(unittest.TestCase):
    """Media sniffing and the adapter dispatch seam."""

    def setUp(self):
        contract._ADAPTERS.clear()

    def tearDown(self):
        contract._ADAPTERS.clear()

    def test_pdf_and_png_are_sniffed_from_signatures(self):
        base = ROOT / 'data/knowledge_base'
        self.assertEqual(contract.detect_media_type(base / '04_售后服务与退款管理制度.pdf'),
                         'application/pdf')
        self.assertEqual(contract.detect_media_type(base / 'f_yoy.png'), 'image/png')
        self.assertEqual(contract.detect_media_type(ROOT / 'requirements.txt'), contract.OCTET_STREAM)

    def test_unreadable_path_is_reported_not_guessed(self):
        with self.assertRaises(HTTPException) as caught:
            contract.detect_media_type(ROOT / 'data/does-not-exist.pdf')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_MEDIA_UNREADABLE')

    def test_dispatch_routes_to_the_registered_handler_and_validates(self):
        seen = {}

        def handler(**kwargs):
            seen.update(kwargs)
            return parse_result()

        contract.register_adapter('demo', media_types=['application/pdf'],
                                  capabilities=['native_text'], handler=handler)
        result = contract.dispatch('application/pdf', marker='x')
        self.assertEqual(seen, {'marker': 'x'})
        self.assertEqual(result.parse_status, 'parsed')
        self.assertEqual(len(result.blocks), 2)

    def test_dispatch_refuses_unknown_media_and_handlerless_adapters(self):
        with self.assertRaises(HTTPException) as caught:
            contract.dispatch('image/tiff')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_ADAPTER_UNKNOWN')
        contract.register_adapter('declared-only', media_types=['application/pdf'],
                                  capabilities=['native_text'])
        with self.assertRaises(HTTPException) as caught:
            contract.dispatch('application/pdf')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_ADAPTER_UNSUPPORTED')

    def test_dispatch_validates_the_handler_output(self):
        contract.register_adapter('bad', media_types=['application/pdf'], capabilities=['x'],
                                  handler=lambda **_: {'classification': {}, 'blocks': []})
        with self.assertRaises(HTTPException) as caught:
            contract.dispatch('application/pdf')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_RESULT_INVALID')

    def test_handler_errors_are_mapped_to_stable_codes(self):
        def rejected(**_):
            raise HTTPException(422, detail={'error_code': 'PARSE_LOCATOR_INVALID'})

        contract.register_adapter('rejecting', media_types=['application/pdf'], capabilities=['x'],
                                  handler=rejected)
        with self.assertRaises(HTTPException) as caught:
            contract.dispatch('application/pdf')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_LOCATOR_INVALID')

        contract._ADAPTERS.clear()
        contract.register_adapter('wrong-signature', media_types=['application/pdf'], capabilities=['x'],
                                  handler=lambda required_positional: parse_result())
        with self.assertRaises(HTTPException) as caught:
            contract.dispatch('application/pdf', unexpected=True)
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_RESULT_INVALID')

        contract._ADAPTERS.clear()
        contract.register_adapter('exploding', media_types=['application/pdf'], capabilities=['x'],
                                  handler=lambda **_: (_ for _ in ()).throw(RuntimeError('boom')))
        with self.assertRaises(HTTPException) as caught:
            contract.dispatch('application/pdf')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_ADAPTER_UNSUPPORTED')

    def test_handler_must_be_callable(self):
        with self.assertRaises(HTTPException) as caught:
            contract.register_adapter('bad-handler', media_types=['application/pdf'],
                                      capabilities=['x'], handler='not-callable')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_RESULT_INVALID')


class StorageColumnTests(unittest.TestCase):
    """The validated locator must reach the block table, and history must stay NULL."""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.patch = patch.object(ingest, 'PAGE_IMG_DIR', Path(self.directory.name))
        self.patch.start()
        self.addCleanup(self.patch.stop)

    def native_pdf(self):
        return ROOT / 'data/knowledge_base/04_售后服务与退款管理制度.pdf'

    def test_parsed_blocks_carry_a_validated_locator_and_version(self):
        blocks, _, detail = ingest.parse_pdf(self.native_pdf(), 'store-native', include_quality=True)
        self.assertTrue(blocks)
        for block in blocks:
            locator = block['contract_locator']
            self.assertIsNotNone(locator, block['block_id'])
            self.assertEqual(locator['locator_kind'], 'pdf_region')
            self.assertEqual(locator['page_no'], block['page_no'])
            self.assertEqual(locator['bbox'], [float(v) for v in block['bbox']])
            self.assertEqual(block['contract_version'], contract.CONTRACT_VERSION)
        self.assertEqual(detail['parse_contract']['locators_written'], len(blocks))

    def locator_block(self, block_id='blk-x'):
        return {'block_id': block_id, 'page_no': 1, 'bbox': [1.0, 2.0, 3.0, 4.0],
                'block_type': 'paragraph', 'text': '正文', 'table_cells': None,
                'formula_latex': None, 'reading_order': 0, 'heading_level': None,
                'ocr_confidence': None, 'contract_version': contract.CONTRACT_VERSION,
                'contract_locator': contract.validate_locator(locator()).payload}

    def test_block_row_keeps_the_historical_shape(self):
        row = ingest._block_row(self.locator_block(), 'doc-1')
        self.assertEqual(len(row), 13)
        self.assertEqual(ingest.BLOCK_INSERT_SQL.count('%s'), 13)
        self.assertNotIn('locator', ingest.BLOCK_INSERT_SQL)

    def test_locator_row_carries_version_and_json_locator(self):
        row = ingest._locator_row(self.locator_block(), 'doc-1')
        self.assertEqual(len(row), 4)
        self.assertEqual(row[1], 'doc-1')
        self.assertEqual(row[2], contract.CONTRACT_VERSION)
        self.assertEqual(json.loads(row[3])['locator_kind'], 'pdf_region')
        self.assertEqual(ingest.LOCATOR_INSERT_SQL.count('%s'), 4)
        self.assertIn('document_block_locators', ingest.LOCATOR_INSERT_SQL)

    def test_block_without_a_locator_writes_no_row(self):
        block = self.locator_block()
        block.pop('contract_locator')
        block.pop('contract_version')
        self.assertIsNone(ingest._locator_row(block, 'doc-1'), '不得为历史块编造 locator')

    def test_statements_actually_fit_the_live_schema(self):
        """Insert the real rows inside a rolled-back transaction: proves column match."""
        import psycopg2
        from datamind.agent import _write_dsn
        block = self.locator_block('blk-schema-check')
        with psycopg2.connect(_write_dsn()) as conn, conn.cursor() as cur:
            cur.execute("""INSERT INTO document_assets(document_id, title, file_type, sha256)
                VALUES ('doc-schema-check','schema check','pdf',%s) ON CONFLICT DO NOTHING""",
                        ('c' * 64,))
            cur.execute(ingest.BLOCK_INSERT_SQL, ingest._block_row(block, 'doc-schema-check'))
            cur.execute(ingest.LOCATOR_INSERT_SQL, ingest._locator_row(block, 'doc-schema-check'))
            cur.execute("""SELECT contract_version, locator->>'locator_kind'
                FROM document_block_locators WHERE block_id=%s""", (block['block_id'],))
            row = cur.fetchone()
            conn.rollback()
        self.assertEqual(row, (contract.CONTRACT_VERSION, 'pdf_region'))

    def test_migration_created_the_locator_table(self):
        import psycopg2
        from datamind.agent import _write_dsn
        with psycopg2.connect(_write_dsn()) as conn, conn.cursor() as cur:
            cur.execute("""SELECT column_name, data_type FROM information_schema.columns
                WHERE table_name='document_block_locators'""")
            found = dict(cur.fetchall())
            cur.execute("""SELECT tableowner FROM pg_tables
                WHERE tablename='document_block_locators'""")
            owner = cur.fetchone()
        self.assertEqual(found.get('contract_version'), 'character varying')
        self.assertEqual(found.get('locator'), 'jsonb')
        self.assertEqual(found.get('block_id'), 'character varying')
        self.assertIsNotNone(owner)


class BlockWidthTests(unittest.TestCase):
    """F06: a value longer than its storage column is refused by the contract.

    The defect: `from_pdf_blocks` accepted a 66-character `block_id` and validation accepted
    a 132-character `document_id`, while every contract table declares VARCHAR(64) for both.
    A result could therefore pass validation and then fail at INSERT time.
    """

    def error(self, code, reason, action):
        with self.assertRaises(HTTPException) as caught:
            action()
        self.assertEqual(caught.exception.detail['error_code'], code)
        self.assertEqual(caught.exception.detail['details'].get('reason'), reason)

    def block(self, document_id='doc-1', **overrides):
        payload = {'block_id': 'b1', 'document_id': document_id, 'source_version': 'v1',
                   'parser_version': 'pdf-ocr-0.4', 'block_type': 'paragraph', 'text': '正文',
                   'native_attributes': {}, 'reading_order': 0,
                   'locator': locator('office_object', document_id=document_id),
                   'quality': {}, 'content_hash': 'a' * 64}
        payload.update(overrides)
        return payload

    def test_the_widths_are_the_storage_widths(self):
        self.assertEqual(contract.BLOCK_COLUMN_WIDTHS,
                         {'block_id': 64, 'document_id': 64, 'parser_version': 64,
                          'source_version': 128})
        self.assertEqual(contract.DOCUMENT_ID_MAX_LENGTH, 64)

    def test_every_identifier_is_checked_against_its_column(self):
        for field, limit in contract.BLOCK_COLUMN_WIDTHS.items():
            if field == 'document_id':
                accepted = self.block(document_id='x' * limit)

                def refused(limit=limit):
                    return contract.validate_block(self.block(document_id='x' * (limit + 1)))
            else:
                accepted = self.block(**{field: 'x' * limit})

                def refused(field=field, limit=limit):
                    return contract.validate_block(self.block(**{field: 'x' * (limit + 1)}))
            self.assertEqual(len(contract.validate_block(accepted).payload[field]), limit)
            self.error('PARSE_BLOCK_INVALID', f'{field}_too_long', refused)

    def test_a_locator_document_id_is_bounded_too(self):
        self.error('PARSE_LOCATOR_INVALID', 'document_id_too_long',
                   lambda: contract.validate_locator(
                       locator('office_object', document_id='d' * 65)))
        self.assertEqual(contract.validate_locator(
            locator('office_object', document_id='d' * 64)).document_id, 'd' * 64)
    def test_require_document_id_is_the_rule_the_adapters_share(self):
        self.assertEqual(contract.require_document_id('d' * 64), 'd' * 64)
        self.error('PARSE_RESULT_INVALID', 'document_id_too_long',
                   lambda: contract.require_document_id('d' * 65, code='PARSE_RESULT_INVALID',
                                                        subject='cad'))

    def generated_pdf_blocks(self, document_id, count=30, **overrides):
        source = {'page_no': 1, 'bbox': [1.0, 2.0, 3.0, 4.0], 'block_type': 'paragraph',
                  'text': '正文'}
        source.update(overrides)
        return contract.from_pdf_blocks([dict(source) for _ in range(count)],
                                        page_sizes={1: (595.0, 842.0)},
                                        parser_version='pdf-ocr-0.4',
                                        document_id=document_id, source_version='v1')

    def test_generated_pdf_block_ids_fit_the_column_for_a_long_document(self):
        mapped = self.generated_pdf_blocks('d' * 64)
        identifiers = [item['block_id'] for item in mapped]
        self.assertEqual(len(set(identifiers)), len(identifiers))
        for identifier in identifiers:
            self.assertLessEqual(len(identifier), 64, identifier)
        self.assertEqual(len(contract.validate_parse_result(
            parse_result(blocks=mapped, relations=[])).blocks), len(mapped))

    def test_an_explicit_over_long_id_is_refused_not_truncated(self):
        # The defect's reproduction: 66 characters used to be accepted here.
        self.error('PARSE_BLOCK_INVALID', 'block_id_too_long',
                   lambda: self.generated_pdf_blocks('doc-1', count=1, block_id='x' * 66))

    def test_a_document_id_the_stores_cannot_hold_is_refused(self):
        for length in (65, 132, 200):
            self.error('PARSE_BLOCK_INVALID', 'document_id_too_long',
                       lambda length=length: self.generated_pdf_blocks('d' * length, count=1))


class BlockShapeClosureTests(unittest.TestCase):
    """F07: every key a validated block may carry has a column to live in.

    Reproducing the defect: validation accepted `heading_level`/`table_cells` while
    `document_contract_blocks` had no column for them, so an Office table lost its cells on
    write even though the contract had said the block was valid.
    """

    def block(self, **overrides):
        payload = {'block_id': 'b1', 'document_id': 'doc-1', 'source_version': 'v1',
                   'parser_version': 'office-native-v1', 'block_type': 'paragraph',
                   'text': '正文', 'native_attributes': {}, 'reading_order': 0,
                   'locator': locator('office_object'), 'quality': {},
                   'content_hash': 'a' * 64}
        payload.update(overrides)
        return payload

    def test_the_block_shape_is_closed(self):
        for key in contract.BLOCK_OPTIONAL_FIELDS:
            self.assertEqual(contract.validate_block(self.block(**{key: None})).payload[key], None)
        with self.assertRaises(HTTPException) as caught:
            contract.validate_block(self.block(invented_field=1))
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_BLOCK_INVALID')
        self.assertEqual(caught.exception.detail['details']['reason'], 'unknown_field')

    def test_the_migrations_declare_a_column_for_every_block_key(self):
        columns = contract_block_columns()
        allowed = set(contract.BLOCK_REQUIRED_FIELDS) | set(contract.BLOCK_OPTIONAL_FIELDS)
        self.assertTrue(allowed <= columns, sorted(allowed - columns))
        # The table may carry its own bookkeeping columns; nothing else is expected.
        self.assertTrue(columns - allowed <= {'contract_version', 'created_at'},
                        sorted(columns - allowed - {'contract_version', 'created_at'}))

    def test_an_empty_key_value_is_still_storable(self):
        """Optional keys that are present but empty must not become "unknown"."""
        block = contract.validate_block(self.block(heading_level=None, table_cells=None)).payload
        self.assertIn('table_cells', block)
        self.assertEqual(block['table_cells'], None)


class OfficeContractBlockStorageTests(unittest.TestCase):
    """F07 acceptance: an Office block round-trips with heading level and table cells."""

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'shape.docx'
        import docx
        document = docx.Document()
        document.add_heading('标题一', level=1)
        document.add_paragraph('正文段落')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '项目'
        table.cell(0, 1).text = '数值'
        table.cell(1, 0).text = '不良率'
        table.cell(1, 1).text = '0.8%'
        document.save(str(self.path))
        self.document_id = f'doc-shape-{uuid.uuid4().hex[:12]}'
        from datamind.agent import _write_dsn
        self.dsn = _write_dsn()
        self.addCleanup(self.cleanup)

    def cleanup(self):
        import psycopg2
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_contract_blocks WHERE document_id=%s',
                        (self.document_id,))

    def parse(self):
        from ingestion import office
        return office.parse(self.path, document_id=self.document_id, source_version='v1')

    def missing_columns(self):
        allowed = set(contract.BLOCK_REQUIRED_FIELDS) | set(contract.BLOCK_OPTIONAL_FIELDS)
        return sorted(allowed - live_contract_block_columns())

    def test_office_blocks_with_headings_and_cells_survive_the_store(self):
        missing = self.missing_columns()
        if missing:
            self.skipTest('db/migrations/017_contract_block_shape.sql not applied: %s' % missing)
        import psycopg2
        from ingestion import contract_store
        payload = self.parse().payload
        title = next(block for block in payload['blocks'] if block['block_type'] == 'title')
        table = next(block for block in payload['blocks'] if block['block_type'] == 'table')
        self.assertEqual(title['heading_level'], 1)
        self.assertEqual(table['table_cells'], [['项目', '数值'], ['不良率', '0.8%']])
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            stored = contract_store.persist_result(cur, document_id=self.document_id,
                                                   payload=payload, dsn=self.dsn)
        self.assertEqual(stored['omitted_columns'], [])
        self.assertEqual(stored['blocks'], len(payload['blocks']))
        read = contract_store.read_result(self.document_id, self.dsn)
        by_type = {block['block_type']: block for block in read['blocks']}
        self.assertEqual(by_type['title']['heading_level'], 1)
        self.assertEqual(by_type['table']['table_cells'], [['项目', '数值'], ['不良率', '0.8%']])
        for block in read['blocks']:
            for name in contract.BLOCK_REQUIRED_FIELDS + contract.BLOCK_OPTIONAL_FIELDS:
                self.assertIn(name, block, name)

    def test_the_live_table_accepts_every_block_key(self):
        missing = self.missing_columns()
        if missing:
            self.skipTest('db/migrations/017_contract_block_shape.sql not applied: %s' % missing)
        import psycopg2
        from psycopg2.extras import Json
        columns = live_contract_block_columns()
        blocks = self.parse().payload['blocks']
        jsonish = ('locator', 'native_attributes', 'quality', 'table_cells', 'bbox')
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            for block in blocks:
                # Every key this particular block carries must have a column; a
                # page-independent Office block simply has no page_no/bbox to store.
                storable = [name for name in (contract.BLOCK_REQUIRED_FIELDS
                                              + contract.BLOCK_OPTIONAL_FIELDS)
                            if name in columns and name in block]
                cur.execute(
                    'INSERT INTO document_contract_blocks(%s) VALUES (%s)' % (
                        ', '.join(['contract_version'] + storable),
                        ', '.join(['%s'] * (1 + len(storable)))),
                    [contract.CONTRACT_VERSION] + [
                        Json(block[name]) if name in jsonish else block[name]
                        for name in storable])
                cur.execute(
                    'SELECT %s FROM document_contract_blocks WHERE block_id=%%s' % ', '.join(storable),
                    (block['block_id'],))
                row = cur.fetchone()
                for name, value in zip(storable, row):
                    self.assertEqual(value, block[name], name)
            conn.rollback()


class CadLocatorProjectionTests(unittest.TestCase):
    """F10: the unified `cad_entity` locator is the projection of the frozen CAD locator."""

    def test_the_declared_source_fields_are_the_frozen_locators_fields(self):
        from ingestion import cad_contract
        self.assertEqual(set(contract.CAD_SOURCE_LOCATOR_FIELDS),
                         set(cad_contract.Locator.model_fields))

    def test_no_source_field_is_declared_dropped(self):
        self.assertEqual(contract.CAD_SOURCE_LOCATOR_DROPPED_FIELDS, {})
        self.assertTrue(set(contract.CAD_SOURCE_LOCATOR_DROPPED_FIELDS)
                        <= set(contract.CAD_SOURCE_LOCATOR_FIELDS))

    def test_every_source_field_is_required_or_optional_for_cad_entity(self):
        covered = set(contract.LOCATOR_REQUIRED_FIELDS['cad_entity']) \
            | set(contract.LOCATOR_OPTIONAL_FIELDS['cad_entity'])
        uncovered = sorted(set(contract.CAD_SOURCE_LOCATOR_FIELDS) - covered)
        self.assertEqual(uncovered, [])

    def test_the_source_digest_is_validated_when_it_is_carried(self):
        carried = contract.validate_locator(locator('cad_entity', source_sha256='a' * 64))
        self.assertEqual(carried.payload['source_sha256'], 'a' * 64)
        with self.assertRaises(HTTPException) as caught:
            contract.validate_locator(locator('cad_entity', source_sha256='not-a-digest'))
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_LOCATOR_INVALID')
        self.assertEqual(caught.exception.detail['details']['reason'], 'source_sha256')

    def test_the_v2_source_fields_are_the_v1_fields_plus_the_block_path(self):
        from ingestion import cad_contract_v2
        self.assertEqual(set(contract.CAD_SOURCE_LOCATOR_V2_FIELDS),
                         set(contract.CAD_SOURCE_LOCATOR_FIELDS) | {'block_path'})
        # The positional identity is the second declared extension of the v2 locator: a
        # handle-less source names its entity by `entity_index` + `handle_source` instead
        # of a handle. Together the two extensions are exactly the v2 source locator.
        self.assertEqual(set(contract.CAD_POSITIONAL_LOCATOR_FIELDS),
                         {'entity_index', 'handle_source'})
        self.assertEqual(set(contract.CAD_SOURCE_LOCATOR_V2_FIELDS)
                         | set(contract.CAD_POSITIONAL_LOCATOR_FIELDS),
                         set(cad_contract_v2.Locator.model_fields))
        covered = set(contract.LOCATOR_REQUIRED_FIELDS['cad_entity']) \
            | set(contract.LOCATOR_OPTIONAL_FIELDS['cad_entity'])
        self.assertEqual(sorted((set(contract.CAD_SOURCE_LOCATOR_V2_FIELDS)
                                 | set(contract.CAD_POSITIONAL_LOCATOR_FIELDS)) - covered), [])

    def test_the_v2_block_path_is_validated_when_it_is_carried(self):
        carried = contract.validate_locator(locator('cad_entity', block_path=['2AF', '3B']))
        self.assertEqual(carried.payload['block_path'], ['2AF', '3B'])
        # An absent chain stays valid: a Model-space entity has no INSERT above it.
        self.assertNotIn('block_path', contract.validate_locator(locator('cad_entity')).payload)

    def test_an_unusable_block_path_is_refused(self):
        for value in ('2AF', ['not-a-handle'], ['2AF', '2AF'], ['2AF'] * 33, [2]):
            with self.subTest(value=value), self.assertRaises(HTTPException) as caught:
                contract.validate_locator(locator('cad_entity', block_path=value))
            self.assertEqual(caught.exception.detail['error_code'], 'PARSE_LOCATOR_INVALID')
            self.assertEqual(caught.exception.detail['details']['subject'], 'locator.cad_entity')

