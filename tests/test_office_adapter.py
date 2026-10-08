"""Office adapters: DOCX/PPTX/XLSX into the unified parse contract, routed by content."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from fastapi import HTTPException

from datamind import parse_contract as contract
from ingestion import office

ROOT = Path(__file__).resolve().parents[1]


def _office_libraries_available():
    try:
        import docx  # noqa: F401
        import openpyxl  # noqa: F401
        import pptx  # noqa: F401
    except ImportError:
        return False
    return True


OFFICE_LIBRARIES = _office_libraries_available()


def docx_fixture(path):
    import docx
    document = docx.Document()
    document.add_heading('第一章 总则', level=1)
    document.add_paragraph('退款需在七个工作日内提出。')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '项目'
    table.cell(0, 1).text = '数值'
    table.cell(1, 0).text = '不良率'
    table.cell(1, 1).text = '0.8%'
    document.save(str(path))


def pptx_fixture(path):
    import pptx
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = '质量回顾'
    slide.placeholders[1].text = '本季度不良率下降'
    presentation.save(str(path))


def xlsx_fixture(path):
    import openpyxl
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '批次表'
    sheet.append(['批次', '温度'])
    sheet.append(['B001', 850])
    sheet.append([None, None])
    sheet.append(['B006', 870])
    workbook.save(str(path))


@unittest.skipUnless(OFFICE_LIBRARIES, "可选 Office 库未安装（requirements-office.txt）")
class OfficeDetectionTests(unittest.TestCase):
    def test_ooxml_subtypes_are_read_from_the_package_index(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            docx, pptx, xlsx = base / 'a.docx', base / 'b.pptx', base / 'c.xlsx'
            docx_fixture(docx)
            pptx_fixture(pptx)
            xlsx_fixture(xlsx)
            self.assertEqual(contract.detect_media_type(docx), office.DOCX)
            self.assertEqual(contract.detect_media_type(pptx), office.PPTX)
            self.assertEqual(contract.detect_media_type(xlsx), office.XLSX)

    def test_a_plain_zip_stays_a_zip_and_is_refused(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'plain.zip'
            with zipfile.ZipFile(path, 'w') as archive:
                archive.writestr('notes.txt', 'hello')
            self.assertEqual(contract.detect_media_type(path), 'application/zip')
            with self.assertRaises(HTTPException) as caught:
                office.parse(path, document_id='doc-1', source_version='v1')
            self.assertEqual(caught.exception.detail['error_code'], 'PARSE_ADAPTER_UNKNOWN')

    def test_missing_library_is_an_explicit_refusal(self):
        with self.assertRaises(HTTPException) as caught:
            office._require('module_that_does_not_exist', 'python-nonexistent')
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_ADAPTER_UNSUPPORTED')


@unittest.skipUnless(OFFICE_LIBRARIES, "可选 Office 库未安装（requirements-office.txt）")
class DocxAdapterTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'spec.docx'
        docx_fixture(self.path)
        office.register_office_adapters()

    def test_paragraphs_headings_and_tables_become_contract_blocks(self):
        result = office.parse(self.path, document_id='doc-docx', source_version='v1')
        blocks = result.payload['blocks']
        self.assertEqual(result.payload['parse_status'], 'parsed')
        self.assertEqual(len(blocks), 3, [b['text'] for b in blocks])
        heading, paragraph, table = blocks
        self.assertEqual(heading['block_type'], 'title')
        self.assertEqual(heading['heading_level'], 1)
        self.assertEqual(heading['text'], '第一章 总则')
        self.assertEqual(paragraph['block_type'], 'paragraph')
        self.assertEqual(paragraph['text'], '退款需在七个工作日内提出。')
        self.assertEqual(table['block_type'], 'table')
        self.assertEqual(table['table_cells'], [['项目', '数值'], ['不良率', '0.8%']])
        for block in blocks:
            locator = block['locator']
            self.assertEqual(locator['locator_kind'], 'office_object')
            self.assertEqual(locator['document_id'], 'doc-docx')
            self.assertEqual(locator['part'], 'word/document.xml')
        self.assertEqual(heading['locator']['paragraph'], 0)
        self.assertEqual(table['locator']['table'], 0)

    def test_every_block_survives_a_full_contract_validation(self):
        result = office.parse(self.path, document_id='doc-docx', source_version='v1')
        validated = contract.validate_parse_result(result.payload)
        self.assertEqual(len(validated.payload['blocks']), 3)

    def test_truncation_is_declared_not_silent(self):
        with patch.object(office, 'MAX_BLOCKS', 2):
            result = office.parse(self.path, document_id='doc-docx', source_version='v1')
        self.assertEqual(result.payload['parse_status'], 'partial')
        self.assertIn('block_limit_reached', result.payload['unsupported_features'])
        self.assertEqual(len(result.payload['blocks']), 2)


@unittest.skipUnless(OFFICE_LIBRARIES, "可选 Office 库未安装（requirements-office.txt）")
class PptxAdapterTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'review.pptx'
        pptx_fixture(self.path)
        office.register_office_adapters()

    def test_slide_text_frames_become_blocks_with_slide_locators(self):
        result = office.parse(self.path, document_id='doc-pptx', source_version='v1')
        blocks = result.payload['blocks']
        self.assertEqual([b['text'] for b in blocks], ['质量回顾', '本季度不良率下降'])
        self.assertEqual(blocks[0]['block_type'], 'title')
        self.assertEqual(blocks[1]['block_type'], 'paragraph')
        for block in blocks:
            self.assertEqual(block['locator']['locator_kind'], 'office_object')
            self.assertEqual(block['locator']['slide'], 1)
            self.assertTrue(block['locator']['part'].startswith('ppt/slides/slide'))
        self.assertEqual(blocks[0]['locator']['shape'], 0)


@unittest.skipUnless(OFFICE_LIBRARIES, "可选 Office 库未安装（requirements-office.txt）")
class XlsxAdapterTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / 'batches.xlsx'
        xlsx_fixture(self.path)
        office.register_office_adapters()

    def test_rows_become_ranged_blocks_and_blank_rows_are_skipped(self):
        result = office.parse(self.path, document_id='doc-xlsx', source_version='v1')
        blocks = result.payload['blocks']
        self.assertEqual([b['text'] for b in blocks], ['批次 | 温度', 'B001 | 850', 'B006 | 870'])
        # Ranges keep the real spreadsheet row numbers: the blank row 3 is skipped, not
        # renumbered, so a citation still points at the row an operator would see.
        self.assertEqual([block['locator']['range'] for block in blocks],
                         ['A1:B1', 'A2:B2', 'A4:B4'])
        for block in blocks:
            locator = block['locator']
            self.assertEqual(locator['locator_kind'], 'sheet_range')
            self.assertEqual(locator['sheet'], '批次表')
            self.assertEqual(block['block_type'], 'table')
        self.assertEqual(blocks[1]['table_cells'], [['B001', '850']])

    def test_row_limit_is_declared_instead_of_truncating_silently(self):
        result = office.parse(self.path, document_id='doc-xlsx', source_version='v1', max_rows_per_sheet=1)
        self.assertEqual(result.payload['parse_status'], 'partial')
        self.assertIn('sheet_row_limit_reached', result.payload['unsupported_features'])
        self.assertEqual(len(result.payload['blocks']), 1)


@unittest.skipUnless(OFFICE_LIBRARIES, "可选 Office 库未安装（requirements-office.txt）")
class OfficeCliTests(unittest.TestCase):
    def test_cli_reports_the_contract_summary_and_refuses_a_plain_zip(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            docx = base / 'spec.docx'
            docx_fixture(docx)
            output = base / 'contract.json'
            completed = subprocess.run(
                [sys.executable, str(ROOT / 'scripts/parse_office.py'), '--source', str(docx),
                 '--document-id', 'doc-cli', '--output', str(output)],
                cwd=str(ROOT), capture_output=True, text=True, timeout=120)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            summary = json.loads(completed.stdout.split('\n契约结果')[0])
            self.assertEqual(summary['adapter'], 'docx-native')
            self.assertEqual(summary['blocks'], 3)
            self.assertEqual(json.loads(output.read_text(encoding='utf-8'))['blocks'][0]['block_type'], 'title')

            plain = base / 'plain.zip'
            with zipfile.ZipFile(plain, 'w') as archive:
                archive.writestr('notes.txt', 'hello')
            refused = subprocess.run(
                [sys.executable, str(ROOT / 'scripts/parse_office.py'), '--source', str(plain),
                 '--document-id', 'doc-zip'],
                cwd=str(ROOT), capture_output=True, text=True, timeout=120)
            self.assertEqual(refused.returncode, 1)
            self.assertIn('PARSE_ADAPTER_UNKNOWN', refused.stderr)


if __name__ == '__main__':
    unittest.main()
