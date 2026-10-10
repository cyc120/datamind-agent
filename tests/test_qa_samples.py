"""WP9 sample set: frozen manifest, deterministic bytes, answer assertions, searchability.

These tests guard the *evidence* the WP9 card rests on, not the model:

* the manifest is frozen (a changed byte, size, or digest on any sample is a failure);
* regenerating a sample reproduces the frozen SHA-256 (so "frozen" means reproducible);
* every question's assertion can be recomputed from the generated file itself (text
  extraction, exact XLSX cells, exact DOCX table cells) or, for the two image-only scans,
  from the ground truth the deterministic rasteriser rendered;
* the retrievability judgement the report quotes is the repository's one function,
  ``ingestion.upload_pipeline.searchability``, exercised read-only (a truth-table call plus
  the tri-state of every real ``document_assets`` row).

Nothing here writes to the database, calls the upstream model, or regenerates the frozen
files on disk: regeneration happens in memory and is compared against the manifest.
"""
from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SAMPLES_DIR = ROOT / 'data' / 'qa_samples'
MANIFEST_PATH = SAMPLES_DIR / 'manifest.json'
REPORT_PATH = ROOT / 'reports' / 'qa_samples_v1.json'

#: The requirement is "≥10 files, ≥4 formats". These five are the ones the task names.
REQUIRED_FORMATS = ('pdf-native', 'pdf-scanned', 'docx-native', 'xlsx-native', 'pptx-native')
SCAN_FORMATS = ('pdf-scanned',)


def _load_builder():
    spec = importlib.util.spec_from_file_location('wp9_build_samples_under_test',
                                                  SAMPLES_DIR / 'build_samples.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


BUILDER = _load_builder()


def _manifest():
    return json.loads(MANIFEST_PATH.read_text(encoding='utf-8'))


class FrozenManifestTest(unittest.TestCase):
    def test_manifest_matches_every_sample(self):
        ok, problems = BUILDER.check_manifest(MANIFEST_PATH)
        self.assertTrue(ok, 'frozen manifest does not match the files: %s' % problems)
        self.assertEqual(problems, [])

    def test_tampered_sample_is_detected(self):
        """A one-byte change to a copy of a sample must fail the frozen check."""
        with tempfile.TemporaryDirectory() as directory:
            scratch = Path(directory)
            shutil.copy2(MANIFEST_PATH, scratch / 'manifest.json')
            manifest = _manifest()
            first = manifest['samples'][0]['file']
            shutil.copy2(SAMPLES_DIR / first, scratch / first)
            target = scratch / first
            data = bytearray(target.read_bytes())
            data[-1] = (data[-1] + 1) % 256
            target.write_bytes(bytes(data))
            ok, problems = BUILDER.check_manifest(scratch / 'manifest.json')
            self.assertFalse(ok)
            self.assertTrue(any('sha256 mismatch' in problem for problem in problems), problems)
            self.assertTrue(any(first in problem for problem in problems), problems)

    def test_tampered_manifest_digest_is_detected(self):
        with tempfile.TemporaryDirectory() as directory:
            scratch = Path(directory)
            manifest = _manifest()
            for entry in manifest['samples']:
                shutil.copy2(SAMPLES_DIR / entry['file'], scratch / entry['file'])
            manifest['samples'][0]['questions'][0]['answer'] = '被篡改的答案'
            (scratch / 'manifest.json').write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
            ok, problems = BUILDER.check_manifest(scratch / 'manifest.json')
            self.assertFalse(ok)
            self.assertTrue(any('digest mismatch' in problem for problem in problems), problems)

    def test_regeneration_reproduces_the_frozen_bytes(self):
        """Every writer is deterministic: regenerating yields the frozen SHA-256.

        This runs in the *test* interpreter, i.e. a different process (and a different
        ``PYTHONHASHSEED``) than the generator that froze the manifest. That difference is
        load-bearing: it is how the DXF ``CLASSES``/``OBJECTS`` record ordering — which
        follows Python set iteration — was caught, while two generations inside one process
        looked identical.
        """
        manifest = _manifest()
        mismatches = []
        for entry in manifest['samples']:
            data = BUILDER.WRITERS[entry['file']]()
            digest = hashlib.sha256(data).hexdigest()
            if digest != entry['sha256']:
                mismatches.append((entry['file'], digest, entry['sha256']))
        self.assertEqual(mismatches, [], 'regeneration is not byte-deterministic: %s' % mismatches)

    def test_every_sample_declares_its_metadata(self):
        manifest = _manifest()
        for entry in manifest['samples']:
            for field in ('file', 'format', 'suffix', 'size', 'sha256', 'text_layer',
                          'expected_locator_kinds', 'expected_locator_fields',
                          'ground_truth', 'questions'):
                self.assertIn(field, entry, entry.get('file'))
            self.assertEqual(len(entry['sha256']), 64, entry['file'])
            self.assertTrue(entry['expected_locator_kinds'], entry['file'])
            self.assertTrue(entry['expected_locator_fields'], entry['file'])


class SampleSetShapeTest(unittest.TestCase):
    def test_at_least_ten_files_and_four_formats(self):
        manifest = _manifest()
        samples = manifest['samples']
        self.assertGreaterEqual(len(samples), 10)
        formats = {entry['format'] for entry in samples}
        self.assertGreaterEqual(len(formats), 4)
        for required in REQUIRED_FORMATS:
            self.assertIn(required, formats)
        self.assertEqual(manifest['counts']['samples'], len(samples))
        self.assertEqual(manifest['counts']['formats'], sorted(formats))

    def test_every_sample_has_at_least_two_answerable_questions(self):
        manifest = _manifest()
        total = 0
        for entry in manifest['samples']:
            self.assertGreaterEqual(len(entry['questions']), 2, entry['file'])
            for question in entry['questions']:
                self.assertTrue(question['question'].strip(), entry['file'])
                self.assertTrue(str(question['answer']).strip(), question['id'])
                self.assertIn(question['assertion']['kind'],
                              ('exact_text', 'arith_sum', 'arith_diff', 'arith_product',
                               'arith_ratio_percent', 'count_matching', 'count_less_than'))
                self.assertIn(question['checkable_from'],
                              ('file_text', 'file_cell', 'ground_truth_only'))
                total += 1
        self.assertEqual(manifest['counts']['questions'], total)

    def test_locator_expectations_follow_the_parse_contract(self):
        from datamind import parse_contract
        manifest = _manifest()
        expected_kind_by_format = {
            'pdf-native': 'pdf_region', 'pdf-scanned': 'pdf_region',
            'docx-native': 'office_object', 'pptx-native': 'office_object',
            'xlsx-native': 'sheet_range', 'dxf-native': 'cad_entity'}
        for entry in manifest['samples']:
            kind = expected_kind_by_format[entry['format']]
            self.assertEqual(entry['expected_locator_kinds'], [kind], entry['file'])
            self.assertIn(kind, parse_contract.LOCATOR_KINDS)
            # The manifest declares the kind-specific fields; the contract's required list is
            # those plus the three identity fields every locator carries.
            identity = ('locator_kind', 'document_id', 'source_version')
            self.assertEqual(entry['expected_locator_fields'][kind],
                             [field for field in parse_contract.LOCATOR_REQUIRED_FIELDS[kind]
                              if field not in identity], entry['file'])

    def test_scanned_samples_have_no_text_layer_and_native_ones_do(self):
        import pymupdf
        manifest = _manifest()
        scanned = 0
        for entry in manifest['samples']:
            if not entry['format'].startswith('pdf'):
                continue
            path = SAMPLES_DIR / entry['file']
            text = BUILDER.extract_text(path)
            with pymupdf.open(str(path)) as document:
                self.assertEqual(document.page_count, entry['pages'], entry['file'])
            if entry['format'] in SCAN_FORMATS:
                scanned += 1
                self.assertFalse(entry['text_layer'], entry['file'])
                self.assertEqual(text.strip(), '', entry['file'])
            else:
                self.assertTrue(entry['text_layer'], entry['file'])
                self.assertTrue(text.strip(), entry['file'])
        self.assertEqual(scanned, 2)


class AnswerAssertionTest(unittest.TestCase):
    def test_assertions_recompute_from_the_files(self):
        manifest = _manifest()
        failures = []
        for entry in manifest['samples']:
            path = SAMPLES_DIR / entry['file']
            for question in entry['questions']:
                check = BUILDER.check_assertion(path, entry, question)
                if not check.get('ok'):
                    failures.append((entry['file'], question['id'], 'assertion', check))
                if not check.get('operand_sources_ok', True):
                    failures.append((entry['file'], question['id'], 'operands', check))
        self.assertEqual(failures, [], failures)

    def test_arithmetic_expectations_hold(self):
        manifest = _manifest()
        arithmetic = 0
        for entry in manifest['samples']:
            for question in entry['questions']:
                assertion = question['assertion']
                kind = assertion['kind']
                if kind == 'arith_sum':
                    computed = sum(assertion['operands'])
                elif kind == 'arith_diff':
                    computed = assertion['operands'][0] - sum(assertion['operands'][1:])
                elif kind == 'arith_product':
                    computed = 1
                    for value in assertion['operands']:
                        computed *= value
                elif kind == 'arith_ratio_percent':
                    computed = round(assertion['operands'][0] / assertion['denominator'] * 100)
                else:
                    continue
                arithmetic += 1
                self.assertEqual(computed, assertion['expected'],
                                 '%s %s' % (entry['file'], question['id']))
        self.assertGreaterEqual(arithmetic, 8)

    def test_xlsx_operands_are_the_real_cells(self):
        """An arithmetic operand declared with a ``cell:`` pointer must equal that cell."""
        import openpyxl
        manifest = _manifest()
        checked = 0
        for entry in manifest['samples']:
            if entry['suffix'] != '.xlsx':
                continue
            path = SAMPLES_DIR / entry['file']
            workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
            try:
                for question in entry['questions']:
                    for source in question['assertion'].get('operand_sources') or []:
                        if not source['pointer'].startswith('cell:'):
                            continue
                        _, sheet, coordinate = source['pointer'].split(':', 2)
                        self.assertEqual(workbook[sheet][coordinate].value, source['value'],
                                         '%s %s' % (entry['file'], source['pointer']))
                        checked += 1
            finally:
                workbook.close()
        self.assertGreaterEqual(checked, 10)

    def test_docx_table_operands_are_the_real_cells(self):
        from docx import Document
        manifest = _manifest()
        checked = 0
        for entry in manifest['samples']:
            if entry['suffix'] != '.docx':
                continue
            document = Document(str(SAMPLES_DIR / entry['file']))
            for question in entry['questions']:
                for source in question['assertion'].get('operand_sources') or []:
                    if not source['pointer'].startswith('docx_table:'):
                        continue
                    _, table, row, column = source['pointer'].split(':')
                    value = document.tables[int(table)].rows[int(row)].cells[int(column)].text
                    self.assertEqual(value, str(source['value']),
                                     '%s %s' % (entry['file'], source['pointer']))
                    checked += 1
        self.assertGreaterEqual(checked, 3)


class SearchabilitySemanticsTest(unittest.TestCase):
    """The report quotes `searchability`; these assertions pin the semantics it means."""

    def test_truth_table_matches_the_upload_pipeline_function(self):
        from ingestion import upload_pipeline
        cases = [
            ('indexed', False, 5, True, upload_pipeline.INDEX_STATE_INDEXED, None),
            ('indexed', False, 0, False, upload_pipeline.INDEX_STATE_NO_CHUNKS,
             upload_pipeline.INDEX_NOTE_NO_CHUNKS),
            ('parsed', False, 0, False, upload_pipeline.INDEX_STATE_NO_CHUNKS,
             upload_pipeline.INDEX_NOTE_NO_CONTRACT_BLOCKS),
            ('parsed', False, 3, False, upload_pipeline.INDEX_STATE_NO_CHUNKS,
             upload_pipeline.INDEX_NOTE_STATUS_NOT_INDEXED),
            ('indexed', True, 5, False, upload_pipeline.INDEX_STATE_NOT_APPROVED,
             upload_pipeline.INDEX_NOTE_NOT_APPROVED),
        ]
        for status, needs_review, chunks, searchable, state, note in cases:
            block_count = 4 if note == upload_pipeline.INDEX_NOTE_NO_CHUNKS else 0
            observed = upload_pipeline.searchability(status, needs_review, chunks,
                                                     block_count=block_count)
            self.assertEqual(observed, (searchable, state, note),
                             (status, needs_review, chunks, block_count))

    def test_stored_rows_agree_with_the_derived_tri_state(self):
        """Every real document row: the function agrees with the documented predicate."""
        import psycopg2
        from psycopg2.extras import RealDictCursor
        from datamind.agent import _read_dsn
        from ingestion import upload_pipeline
        try:
            connection = psycopg2.connect(_read_dsn())
        except psycopg2.OperationalError as error:  # pragma: no cover - environment dependent
            self.skipTest('metadata database unavailable: %s' % error)
        with connection as conn, conn.cursor(cursor_factory=RealDictCursor) as cur:
            conn.set_session(readonly=True)
            cur.execute('''SELECT a.document_id, a.status, a.needs_review,
                    (SELECT count(*) FROM chunks c WHERE c.document_id=a.document_id) AS chunks,
                    (SELECT count(*) FROM document_contract_blocks b
                      WHERE b.document_id=a.document_id) AS blocks
                FROM document_assets a''')
            rows = [dict(row) for row in cur.fetchall()]
        self.assertGreater(len(rows), 0)
        for row in rows:
            searchable, state, note = upload_pipeline.searchability(
                row['status'], row['needs_review'], row['chunks'], block_count=row['blocks'])
            expected = (row['status'] == 'indexed' and not row['needs_review']
                        and int(row['chunks'] or 0) > 0)
            self.assertEqual(searchable, expected, row['document_id'])
            if searchable:
                self.assertEqual((state, note), (upload_pipeline.INDEX_STATE_INDEXED, None),
                                 row['document_id'])
            else:
                self.assertIn(state, (upload_pipeline.INDEX_STATE_NO_CHUNKS,
                                      upload_pipeline.INDEX_STATE_NOT_APPROVED),
                              row['document_id'])
                self.assertIsNotNone(note, row['document_id'])
            # `approve` never publishes an upload that has no chunks: this is the exact
            # property the WP9 card claims ("无切片不可检索").
            if int(row['chunks'] or 0) == 0:
                self.assertFalse(searchable, row['document_id'])

    def test_report_consistency_section_when_present(self):
        if not REPORT_PATH.is_file():
            self.skipTest('reports/qa_samples_v1.json not generated in this checkout')
        report = json.loads(REPORT_PATH.read_text(encoding='utf-8'))
        stored = report.get('stored') or {}
        consistency = stored.get('consistency') or {}
        self.assertTrue(stored.get('manifest', {}).get('frozen'), 'manifest was not frozen')
        self.assertEqual(consistency.get('inconsistent'), 0, consistency)
        self.assertEqual(consistency.get('consistent'), consistency.get('documents_with_db_row'),
                         consistency)
        self.assertEqual(consistency.get('assertions_failed'), 0, consistency)
        # The honest conclusion must be present: the upload path has no vector index.
        self.assertIn('automatic_index', report.get('index_wiring') or {})
        self.assertEqual(report['index_wiring']['automatic_index'], 'not_done')


if __name__ == '__main__':
    unittest.main()
