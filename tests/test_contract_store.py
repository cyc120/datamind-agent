"""Contract store: persisting validated blocks, locators and relations, and reading back."""
import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
import uuid
from unittest.mock import patch

from fastapi import HTTPException
from fastapi.testclient import TestClient
import psycopg2

from datamind.agent import _write_dsn
from datamind.api import app
from datamind import parse_contract as contract
from ingestion import cad_adapter, contract_store, office

ROOT = Path(__file__).resolve().parents[1]
DXF = ROOT / 'data/cad_demo/development/dev-mm.dxf'
IMPORT_SCRIPT = ROOT / 'scripts/import_contract.py'


def load_import_script():
    """The CLI is a script, not an importable package module; load it by path."""
    spec = importlib.util.spec_from_file_location('import_contract_under_test', IMPORT_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def dxf_payload(document_id='doc-store-check'):
    cad_adapter.register_cad_adapters()
    return cad_adapter.parse(DXF, document_id=document_id).payload


def docx_fixture(path):
    import docx
    document = docx.Document()
    document.add_heading('第一章 总则', level=1)
    document.add_paragraph('退款需在七个工作日内提出。')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '项目'
    table.cell(0, 1).text = '数值'
    table.cell(1, 0).text = '不良率'
    table.cell(1, 1).text = '0.8%'
    document.save(str(path))


class ContractStoreTests(unittest.TestCase):
    def setUp(self):
        self.dsn = _write_dsn()
        self.document_id = 'doc-contract-store-test'
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('''INSERT INTO document_assets(document_id, title, file_type, sha256)
                VALUES (%s, %s, 'dxf', %s) ON CONFLICT (document_id) DO NOTHING''',
                        (self.document_id, 'contract store test', 'c' * 64))
            cur.execute('DELETE FROM document_contract_blocks WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_block_relations WHERE document_id=%s', (self.document_id,))

    def tearDown(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_contract_blocks WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_block_relations WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_assets WHERE document_id=%s', (self.document_id,))

    def persist(self, payload=None):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            return contract_store.persist_result(cur, document_id=self.document_id,
                                                 payload=payload or dxf_payload(self.document_id))

    def test_blocks_locators_and_relations_round_trip(self):
        stored = self.persist()
        self.assertEqual(stored['blocks'], 8)
        self.assertEqual(stored['locators'], 8)
        self.assertEqual(stored['relations'], 1)
        self.assertEqual(stored['parse_status'], 'parsed')

        read = contract_store.read_result(self.document_id, self.dsn)
        self.assertEqual(read['block_count'], 8)
        self.assertEqual(read['relation_count'], 1)
        first = read['blocks'][0]
        self.assertEqual(first['locator']['locator_kind'], 'cad_entity')
        self.assertEqual(first['locator']['unit'], 'mm')
        self.assertEqual(first['contract_version'], contract.CONTRACT_VERSION)
        self.assertTrue(first['content_hash'])
        relation = read['relations'][0]
        self.assertEqual(relation['relation_type'], 'references')
        self.assertEqual(relation['relation_state'], 'verified')
        self.assertEqual(relation['recovery_method'], 'native_reference')
        self.assertTrue(relation['evidence'])
        self.assertNotEqual(relation['source_block_id'], relation['target_block_id'])

    def test_persisting_again_replaces_rather_than_duplicates(self):
        self.persist()
        self.persist()
        read = contract_store.read_result(self.document_id, self.dsn)
        self.assertEqual(read['block_count'], 8)
        self.assertEqual(read['relation_count'], 1)

    def test_relation_ids_are_deterministic(self):
        payload = dxf_payload(self.document_id)
        relation = payload['relations'][0]
        first = contract_store.relation_id(self.document_id, relation)
        second = contract_store.relation_id(self.document_id, relation)
        self.assertEqual(first, second)
        self.assertTrue(first.startswith('rel-'))
        self.assertNotEqual(first, contract_store.relation_id('doc-other', relation))

    def test_an_invalid_result_is_refused_before_writing(self):
        payload = dxf_payload(self.document_id)
        payload['blocks'][0]['locator']['bbox'] = [1, 2, 3, 4]  # cad_entity has no bbox
        with self.assertRaises(HTTPException) as caught:
            self.persist(payload)
        self.assertEqual(caught.exception.detail['error_code'], 'PARSE_LOCATOR_INVALID')
        self.assertEqual(contract_store.read_result(self.document_id, self.dsn)['block_count'], 0)

    def test_missing_evidence_is_refused_by_the_database_constraint(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('''INSERT INTO document_contract_blocks(block_id, document_id,
                contract_version, block_type, text, reading_order, locator, content_hash)
                VALUES ('blk-a', %s, 'parse-contract-v1', 'paragraph', 'x', 0, '{}', %s)''',
                        (self.document_id, 'd' * 64))
            with self.assertRaises(psycopg2.Error):
                cur.execute('''INSERT INTO document_block_relations(relation_id, document_id,
                    source_block_id, target_block_id, relation_type, relation_state,
                    recovery_method, evidence) VALUES ('rel-x', %s, 'blk-a', 'blk-a',
                    'references', 'verified', 'native_reference', '')''', (self.document_id,))
            conn.rollback()


class DeclaredGapTests(unittest.TestCase):
    """A document whose format cannot carry a locator must expose that gap, not look empty.

    An OHR-Bench JSON record has a page index but no page geometry, so no honest
    `pdf_region` exists (inventing `[0,0,0,0]` was the defect). The counts alone then read
    as "indexed but nothing stored", so the endpoint repeats the declared gap.
    """

    def setUp(self):
        self.dsn = _write_dsn()
        self.document_id = f'doc-gap-{uuid.uuid4().hex[:12]}'
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute("""INSERT INTO document_assets(document_id, title, file_type, sha256,
                status, acl, needs_review, collection, quality_detail)
                VALUES (%s, 'gap test', 'json', %s, 'indexed', '{public}', FALSE, 'demo_kb', %s)
                ON CONFLICT (document_id) DO NOTHING""",
                        (self.document_id, 'a' * 64, json.dumps({'parse_contract': {
                            'validation': 'parsed', 'adapter': 'ohr-bench-json',
                            'parse_status': 'unsupported', 'blocks': 0, 'locators_written': 0,
                            'unsupported_features': ['locator_unavailable']}})))
        self.client = TestClient(app)
        self.addCleanup(self.cleanup)

    def cleanup(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_assets WHERE document_id=%s', (self.document_id,))

    def test_the_declared_gap_is_visible_instead_of_reading_as_empty(self):
        response = self.client.get(f'/v1/documents/{self.document_id}/contract')
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body['block_count'], 0)
        self.assertEqual(body['locator_count'], 0)
        self.assertIn('locator_unavailable', body['unsupported_features'])
        self.assertEqual(body['parse_contract']['parse_status'], 'unsupported')


class ContentHashSemanticsTests(unittest.TestCase):
    """The two `content_hash` values in one read result must never be swapped.

    `field_store` verifies evidence with `sha256(text)`; the contract hash is a promise over
    identity + source + position + text. Mixing them makes fact verification fail
    *silently* (it returns no facts instead of raising), so the difference is pinned here.
    """

    def test_the_two_hashes_are_derived_differently(self):
        import hashlib
        text = '退款需在七个工作日内提出。'
        page_hash = hashlib.sha256(text.encode()).hexdigest()
        contract_hash = hashlib.sha256(repr(['doc-1', 'v1', 'parse-contract-v1', 1,
                                             'blk-1', text, [1.0, 2.0, 3.0, 4.0]]).encode()).hexdigest()
        self.assertNotEqual(page_hash, contract_hash)
        self.assertEqual(len(page_hash), 64)
        self.assertEqual(len(contract_hash), 64)

    def test_read_result_documents_the_difference(self):
        # The docstring is the contract for callers; if it is dropped, the ambiguity that
        # caused this finding returns silently.
        source = __import__('inspect').getsource(contract_store.read_result)
        self.assertIn('not', source)
        self.assertIn('sha256(text)', source.replace('`', ''))


class ContractEndpointTests(unittest.TestCase):
    def setUp(self):
        self.dsn = _write_dsn()
        self.document_id = 'doc-contract-api-test'
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('''INSERT INTO document_assets(document_id, title, file_type, sha256,
                status, acl, needs_review, collection)
                VALUES (%s, 'contract api test', 'dxf', %s, 'indexed', '{public}', FALSE, 'demo_kb')
                ON CONFLICT (document_id) DO UPDATE SET status='indexed', acl='{public}',
                needs_review=FALSE''', (self.document_id, 'e' * 64))
            contract_store.persist_result(cur, document_id=self.document_id,
                                          payload=dxf_payload(self.document_id))
        self.client = TestClient(app)
        self.addCleanup(self.cleanup)

    def cleanup(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_block_locators WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_blocks WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_contract_blocks WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_block_relations WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_assets WHERE document_id=%s', (self.document_id,))

    def test_endpoint_also_returns_page_shaped_locators(self):
        """A PDF parse keeps its locator in the other store; the endpoint must show it.

        Regression: the endpoint used to read only `document_contract_blocks`, so a PDF
        with 33 validated locators reported `block_count: 0` — "nothing stored" for a
        document whose evidence was fully located.
        """
        page_block = 'blk-pdf-locator-test'
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute("""INSERT INTO document_blocks(block_id, document_id, page_no, bbox,
                block_type, text, reading_order, parser_version, content_hash)
                VALUES (%s, %s, 1, '[1,2,3,4]', 'paragraph', '正文', 0, 'pdf-ocr-0.4', %s)
                ON CONFLICT (block_id) DO NOTHING""",
                        (page_block, self.document_id, 'f' * 64))
            cur.execute("""INSERT INTO document_block_locators(block_id, document_id,
                contract_version, locator) VALUES (%s, %s, %s, %s)
                ON CONFLICT (block_id) DO UPDATE SET locator=EXCLUDED.locator""",
                        (page_block, self.document_id, contract.CONTRACT_VERSION,
                         json.dumps({'locator_kind': 'pdf_region', 'document_id': self.document_id,
                                     'source_version': 'v1', 'page_no': 1, 'page_width': 595.0,
                                     'page_height': 842.0, 'bbox': [1.0, 2.0, 3.0, 4.0],
                                     'unit': 'point', 'coordinate_system': 'top-left',
                                     'rotation_degrees': 0})))
        self.addCleanup(self.drop_page_block, page_block)
        response = self.client.get(f'/v1/documents/{self.document_id}/contract')
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body['locator_count'], 1)
        self.assertEqual(body['locators'][0]['locator']['locator_kind'], 'pdf_region')
        self.assertEqual(body['locators'][0]['page_no'], 1)

    def drop_page_block(self, block_id):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_block_locators WHERE block_id=%s', (block_id,))
            cur.execute('DELETE FROM document_blocks WHERE block_id=%s', (block_id,))

    def test_endpoint_returns_locators_and_relations(self):
        response = self.client.get(f'/v1/documents/{self.document_id}/contract')
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body['block_count'], 8)
        self.assertEqual(body['relation_count'], 1)
        self.assertEqual(body['blocks'][0]['locator']['locator_kind'], 'cad_entity')
        self.assertEqual(body['relations'][0]['relation_type'], 'references')
        # The payload must be JSON-serializable as-is.
        self.assertIn('cad_entity', json.dumps(body))

    def test_endpoint_is_404_for_an_unknown_document(self):
        response = self.client.get('/v1/documents/doc-does-not-exist/contract')
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()['detail']['error_code'], 'DOCUMENT_NOT_FOUND')


class ImportContractCliTests(unittest.TestCase):
    """The CLI must register the asset row that `GET /v1/documents/{id}/contract` reads first.

    Regression: the CLI wrote only `document_contract_blocks`, so `_public_document()` found
    no `document_assets` row and the endpoint answered 404 `DOCUMENT_NOT_FOUND` even though
    the blocks were stored.
    """

    def setUp(self):
        self.dsn = _write_dsn()
        self.token = uuid.uuid4().hex[:12]
        self.document_id = f'doc-cli-import-{self.token}'
        self.title = f'cli import {self.token}'
        self.collection = 'demo_kb'
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.source = Path(self.tempdir.name) / f'cli-import-{self.token}.docx'
        docx_fixture(self.source)
        self.sha256 = hashlib.sha256(self.source.read_bytes()).hexdigest()
        self.module = load_import_script()
        self.client = TestClient(app)
        self.addCleanup(self.cleanup)

    def cleanup(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_contract_blocks WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_block_relations WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_assets WHERE document_id=%s', (self.document_id,))

    def run_cli(self):
        buffer = io.StringIO()
        with patch.dict(os.environ, {'DATAMIND_META_DSN': self.dsn}):
            with contextlib.redirect_stdout(buffer):
                code = self.module.main(['--source', str(self.source),
                                         '--document-id', self.document_id,
                                         '--collection', self.collection,
                                         '--title', self.title, '--dsn', self.dsn])
        return code, buffer.getvalue()

    def asset_rows(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('''SELECT document_id, title, collection, file_type, sha256, status,
                acl, needs_review, origin FROM document_assets WHERE document_id=%s''',
                        (self.document_id,))
            return cur.fetchall()

    def test_a_single_import_registers_one_asset_row_with_a_visible_status(self):
        code, output = self.run_cli()
        self.assertEqual(code, 0, output)
        summary = json.loads(output)
        self.assertTrue(summary['persisted'])
        self.assertEqual(summary['asset']['file_type'], 'docx')
        self.assertEqual(summary['asset']['status'], 'indexed')
        self.assertEqual(summary['asset']['collection'], self.collection)
        rows = self.asset_rows()
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row[:5], (self.document_id, self.title, self.collection, 'docx', self.sha256))
        # `_public_document()` needs indexed + public + not needs_review + a known collection.
        self.assertEqual(row[5], 'indexed')
        self.assertEqual(list(row[6]), ['public'])
        self.assertFalse(row[7])

    def test_the_contract_endpoint_reads_an_imported_document_instead_of_404(self):
        self.assertEqual(self.run_cli()[0], 0)
        with patch.dict(os.environ, {'DATAMIND_META_DSN': self.dsn}):
            response = self.client.get(f'/v1/documents/{self.document_id}/contract')
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body['block_count'], 3)
        self.assertEqual(body['document_id'], self.document_id)
        self.assertEqual(body['source_sha256'], self.sha256)
        # A DOCX parser keeps its locators inside the contract blocks (`blocks`), not in the
        # page-shaped `locators` list, so only `blocks` can carry them here.
        self.assertEqual(body['locators'], [])
        block_kinds = {block['locator']['locator_kind'] for block in body['blocks']}
        self.assertEqual(block_kinds, {'office_object'})
        self.assertEqual({block['contract_version'] for block in body['blocks']},
                         {contract.CONTRACT_VERSION})
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('SELECT file_type FROM document_assets WHERE document_id=%s', (self.document_id,))
            self.assertEqual(cur.fetchone()[0], 'docx')

    def test_importing_the_same_file_twice_keeps_exactly_one_asset_row(self):
        self.assertEqual(self.run_cli()[0], 0)
        self.assertEqual(self.run_cli()[0], 0)
        self.assertEqual(len(self.asset_rows()), 1)
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('SELECT count(*) FROM document_contract_blocks WHERE document_id=%s',
                        (self.document_id,))
            first = cur.fetchone()[0]
        self.assertEqual(self.run_cli()[0], 0)
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('SELECT count(*) FROM document_contract_blocks WHERE document_id=%s',
                        (self.document_id,))
            self.assertEqual(cur.fetchone()[0], first)
        self.assertEqual(len(self.asset_rows()), 1)

    def test_a_dry_run_writes_no_asset_row(self):
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer):
            code = self.module.main(['--source', str(self.source), '--document-id', self.document_id,
                                     '--collection', self.collection, '--dry-run', '--dsn', self.dsn])
        self.assertEqual(code, 0, buffer.getvalue())
        self.assertEqual(self.asset_rows(), [])


class ReadLocatorsAcrossStoresTests(unittest.TestCase):
    """`read_locators` must cover both stores so Office/DXF locators are not invisible.

    Regression: it queried only `document_block_locators` (the page-shaped PDF store), so a
    DXF/Office block — whose locator lives in `document_contract_blocks.locator` — never
    reached a citation.
    """

    def setUp(self):
        self.dsn = _write_dsn()
        self.document_id = f'doc-locator-union-{uuid.uuid4().hex[:12]}'
        self.addCleanup(self.cleanup)

    def cleanup(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_block_locators WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_contract_blocks WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_block_relations WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_assets WHERE document_id=%s', (self.document_id,))

    def test_office_and_dxf_locators_are_read_from_the_contract_store(self):
        payload = dxf_payload(self.document_id)
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            contract_store.persist_result(cur, document_id=self.document_id, payload=payload)
        block_ids = [block['block_id'] for block in payload['blocks']]
        found = contract_store.read_locators(block_ids, self.dsn)
        self.assertEqual(sorted(found), sorted(block_ids))
        self.assertEqual(found[block_ids[0]]['locator']['locator_kind'], 'cad_entity')
        self.assertEqual(found[block_ids[0]]['contract_version'], contract.CONTRACT_VERSION)

    def test_a_page_shaped_locator_wins_when_one_id_exists_in_both_stores(self):
        payload = dxf_payload(self.document_id)
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            contract_store.persist_result(cur, document_id=self.document_id, payload=payload)
        block_id = payload['blocks'][0]['block_id']
        page_locator = {'locator_kind': 'pdf_region', 'document_id': self.document_id,
                        'source_version': 'v1', 'page_no': 1, 'page_width': 595.0,
                        'page_height': 842.0, 'bbox': [1.0, 2.0, 3.0, 4.0], 'unit': 'point',
                        'coordinate_system': 'top-left', 'rotation_degrees': 0}
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('''INSERT INTO document_block_locators(block_id, document_id,
                contract_version, locator) VALUES (%s, %s, 'parse-contract-v1', %s)
                ON CONFLICT (block_id) DO UPDATE SET locator=EXCLUDED.locator''',
                        (block_id, self.document_id, json.dumps(page_locator)))
        found = contract_store.read_locators([block_id], self.dsn)
        self.assertEqual(list(found), [block_id])
        self.assertEqual(found[block_id]['locator']['locator_kind'], 'pdf_region')
        self.assertEqual(found[block_id]['locator']['bbox'], [1.0, 2.0, 3.0, 4.0])
        self.assertEqual(found[block_id]['contract_version'], 'parse-contract-v1')

    def test_other_ids_and_bounds_behaviour_are_unchanged(self):
        payload = dxf_payload(self.document_id)
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            contract_store.persist_result(cur, document_id=self.document_id, payload=payload)
        block_id = payload['blocks'][0]['block_id']
        found = contract_store.read_locators([block_id, 'blk-not-stored-anywhere'], self.dsn)
        self.assertEqual(list(found), [block_id])
        self.assertEqual(contract_store.read_locators([], self.dsn), {})
        self.assertEqual(contract_store.read_locators([None, ''], self.dsn), {})
        with self.assertRaises(ValueError) as caught:
            contract_store.read_locators(
                [f'blk-{index}' for index in range(contract_store.MAX_LOCATOR_QUERY + 1)], self.dsn)
        self.assertEqual(str(caught.exception), 'CONTRACT_LOCATOR_QUERY_TOO_LARGE')


class ContractBlockShapeTests(unittest.TestCase):
    """Every field the contract validated must reach a column, not just the nine core ones.

    Regression (F07): `document_contract_blocks` had no column for `heading_level`,
    `table_cells`, `parser_version`, … so a block the contract called valid was written
    without those fields and they could never be read back.
    """

    def setUp(self):
        self.dsn = _write_dsn()
        self.document_id = f'doc-block-shape-{uuid.uuid4().hex[:12]}'
        self.source_version = 'a' * 64
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.source = Path(self.tempdir.name) / 'block-shape.docx'
        docx_fixture(self.source)
        office.register_office_adapters()
        self.payload = contract.dispatch(contract.detect_media_type(self.source), path=self.source,
                                         document_id=self.document_id,
                                         source_version=self.source_version).payload
        self.by_type = {block['block_type']: block for block in self.payload['blocks']}
        self.addCleanup(self.cleanup)

    def cleanup(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_contract_blocks WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_block_relations WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_assets WHERE document_id=%s', (self.document_id,))

    def persist(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            return contract_store.persist_result(cur, document_id=self.document_id,
                                                 payload=self.payload, dsn=self.dsn)

    def raw_row(self, block_id, columns):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute(f"SELECT {', '.join(columns)} FROM document_contract_blocks "
                        'WHERE block_id=%s', (block_id,))
            return cur.fetchone()

    def test_optionals_are_written_to_their_own_columns(self):
        stored = self.persist()
        self.assertEqual(stored.get('omitted_columns'), [])
        title = self.by_type['title']
        table = self.by_type['table']
        # Assert against the raw columns, not only the read-back API: the write is what the
        # defect was about, and a self-consistent reader could hide a dropped column.
        title_row = self.raw_row(title['block_id'],
                                 ('source_version', 'parser_version', 'quality', 'heading_level',
                                  'table_cells', 'formula_latex', 'ocr_confidence', 'page_no', 'bbox'))
        self.assertEqual(title_row[0], self.source_version)
        self.assertEqual(title_row[1], title['parser_version'])
        self.assertEqual(title_row[2], title.get('quality'))
        self.assertEqual(title_row[3], 1)
        self.assertIsNone(title_row[4])          # a title block carries no table grid
        self.assertIsNone(title_row[5])          # no formula
        self.assertIsNone(title_row[6])          # no OCR origin
        self.assertIsNone(title_row[7])          # not page-shaped
        self.assertIsNone(title_row[8])          # no bbox
        table_row = self.raw_row(table['block_id'],
                                 ('table_cells', 'heading_level', 'parser_version', 'quality'))
        self.assertEqual(table_row[0], table['table_cells'])
        self.assertIsNone(table_row[1])
        self.assertEqual(table_row[2], table['parser_version'])
        self.assertEqual(table_row[3], table.get('quality'))

    def test_read_result_returns_every_field_unchanged(self):
        self.persist()
        read = contract_store.read_result(self.document_id, self.dsn)
        self.assertEqual(read['block_count'], len(self.payload['blocks']))
        by_type = {block['block_type']: block for block in read['blocks']}
        for block_type, original in self.by_type.items():
            read_back = by_type[block_type]
            for key, value in original.items():
                if key in ('native_attributes', 'locator', 'quality'):
                    continue
                self.assertEqual(read_back[key], value, f'{block_type}.{key}')
            self.assertEqual(read_back['quality'], original.get('quality') or {})
        # Keys the row genuinely has no value for are present but None, never missing.
        for block in read['blocks']:
            for name in contract_store.BLOCK_EXTENDED_COLUMNS:
                self.assertIn(name, block)
        self.assertEqual(by_type['title']['heading_level'], 1)
        self.assertEqual(by_type['table']['table_cells'], self.by_type['table']['table_cells'])
        self.assertEqual(by_type['title']['source_version'], self.source_version)
        self.assertEqual(by_type['title']['parser_version'], self.by_type['title']['parser_version'])

    def test_is_idempotent_still_and_keeps_the_extended_columns(self):
        self.persist()
        self.persist()
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('SELECT count(*) FROM document_contract_blocks WHERE document_id=%s',
                        (self.document_id,))
            self.assertEqual(cur.fetchone()[0], len(self.payload['blocks']))
        read = contract_store.read_result(self.document_id, self.dsn)
        title = next(block for block in read['blocks'] if block['block_type'] == 'title')
        self.assertEqual(title['heading_level'], 1)


class PreMigrationShapeFallbackTests(unittest.TestCase):
    """On a database without the 017 columns the writer must degrade, not lose the block.

    The pre-migration shape is simulated by removing the extended columns from the store's
    own column set, which is exactly what the database probe reports on a database that has
    not applied migration 017.
    """

    def setUp(self):
        self.dsn = _write_dsn()
        self.document_id = f'doc-premigration-{uuid.uuid4().hex[:12]}'
        self.payload = dxf_payload(self.document_id)
        self.addCleanup(self.cleanup)

    def cleanup(self):
        with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
            cur.execute('DELETE FROM document_contract_blocks WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_block_relations WHERE document_id=%s', (self.document_id,))
            cur.execute('DELETE FROM document_assets WHERE document_id=%s', (self.document_id,))

    def test_the_core_columns_are_still_written_and_the_gap_is_reported(self):
        with patch.object(contract_store, 'BLOCK_EXTENDED_COLUMNS', ()):
            with psycopg2.connect(self.dsn) as conn, conn.cursor() as cur:
                stored = contract_store.persist_result(cur, document_id=self.document_id,
                                                       payload=self.payload, dsn=self.dsn)
            read = contract_store.read_result(self.document_id, self.dsn)
        self.assertEqual(stored['blocks'], len(self.payload['blocks']))
        self.assertEqual(read['block_count'], len(self.payload['blocks']))
        first = read['blocks'][0]
        self.assertEqual(first['locator']['locator_kind'], 'cad_entity')
        self.assertTrue(first['content_hash'])
        self.assertEqual(first['reading_order'], 0)
        # Nothing extra may be claimed when the store had nowhere to put it.
        self.assertNotIn('heading_level', first)


if __name__ == '__main__':
    unittest.main()
