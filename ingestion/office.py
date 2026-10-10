"""Office adapters: DOCX / PPTX / XLSX into the unified parse contract.

Each adapter reads the real document object model and emits contract blocks whose
locators point back into the OOXML part (paragraph / table / slide / shape) or the sheet
range, so a citation can name the exact object it came from.

Scope is deliberately bounded: text, headings, table cells and sheet rows only. Styles,
merged-cell semantics, hidden sheets, images and charts are not interpreted, and a document
larger than `MAX_BLOCKS` is truncated with an explicit `unsupported_features` entry instead
of silently losing content.

Two fidelity rules are enforced rather than assumed: DOCX blocks follow the real body order
(`w:p` / `w:tbl` interleaving), and an XLSX formula whose cached value is absent is declared
as `formula_without_cached_value` instead of being reported as a complete read.
"""
from __future__ import annotations

import hashlib
from itertools import zip_longest
from pathlib import Path

from datamind import parse_contract

# v2: block ids are digests, DOCX follows body order, XLSX formula gaps are declared.
CONTRACT_PARSER_VERSION = 'office-adapter-v2'
MAX_BLOCKS = 2000
DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
PPTX = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
ADAPTERS = {'docx-native': DOCX, 'pptx-native': PPTX, 'xlsx-native': XLSX}


def _require(module_name, package):
    try:
        return __import__(module_name)
    except ImportError:
        parse_contract.fail('PARSE_ADAPTER_UNSUPPORTED', subject='office', reason='library_missing',
                            package=package)


def _block_identifier(document_id, index):
    """Return a stable block id that no other (document_id, index) can share.

    The previous scheme built ``off-{document_id}-{index}`` and truncated it at 64
    characters: a long document_id collapsed every block onto the same prefix, so distinct
    blocks reached the contract with duplicate ids. A fixed-length digest of the pair is
    collision-free for any document_id and always fits the 64-character budget.
    """
    material = f'{document_id}\x1f{index}'.encode('utf-8')
    return 'off-' + hashlib.sha256(material).hexdigest()[:32]


def _block(document_id, source_version, block_type, text, locator, *, index, heading_level=None,
           table_cells=None, native_attributes=None, quality=None):
    payload = {'block_id': _block_identifier(document_id, index),
               'document_id': document_id, 'source_version': source_version,
               'parser_version': CONTRACT_PARSER_VERSION, 'block_type': block_type,
               'text': text or '', 'native_attributes': native_attributes or {},
               'reading_order': index, 'locator': locator, 'quality': quality or {},
               'content_hash': _content_hash(document_id, source_version, locator, text, index)}
    if heading_level is not None:
        payload['heading_level'] = heading_level
    if table_cells is not None:
        payload['table_cells'] = table_cells
    return payload


def _content_hash(document_id, source_version, locator, text, index):
    return hashlib.sha256(repr([document_id, source_version, CONTRACT_PARSER_VERSION, index,
                                text or '', locator]).encode('utf-8')).hexdigest()


def _result(document_id, media_type, blocks, *, adapter, parse_status, unsupported=None,
            source_name=None):
    return {'classification': {'detected_media_type': media_type, 'source_name': source_name},
            'adapter_capabilities': [adapter],
            'blocks': blocks, 'relations': [],
            'parsed_scope': [{'kind': 'object', 'ref': block['block_id'], 'status': 'parsed'}
                             for block in blocks],
            'failed_scope': [],
            'quality_findings': [], 'unsupported_features': list(unsupported or []),
            'artifact_refs': [],
            'manifest': {'parser_version': CONTRACT_PARSER_VERSION, 'adapter': adapter},
            'parse_status': parse_status}


def _truncated(counter, unsupported):
    if counter >= MAX_BLOCKS:
        if 'block_limit_reached' not in unsupported:
            unsupported.append('block_limit_reached')
        return True
    return False


def _docx_body_sequence(document):
    """Return the DOCX body objects in true document order, plus an order-trust flag.

    python-docx exposes paragraphs and tables as two separate collections, but the body is an
    ordered sequence of ``w:p`` and ``w:tbl`` children. Walking that sequence keeps the real
    interleaving, while the positions still index the same ``document.paragraphs`` /
    ``document.tables`` entries that the locators name. `complete` is False when the XML body
    and those collections disagree; objects the walk missed are then appended so nothing is
    silently dropped, and the caller declares the unproven ordering.
    """
    from docx.oxml.ns import qn
    paragraphs, tables = list(document.paragraphs), list(document.tables)
    sequence, seen_paragraphs, seen_tables = [], set(), set()
    paragraph_no = table_no = 0
    for child in document.element.body.iterchildren():
        if child.tag == qn('w:p'):
            if paragraph_no < len(paragraphs):
                seen_paragraphs.add(paragraph_no)
                sequence.append(('paragraph', paragraph_no, paragraphs[paragraph_no]))
            paragraph_no += 1
        elif child.tag == qn('w:tbl'):
            if table_no < len(tables):
                seen_tables.add(table_no)
                sequence.append(('table', table_no, tables[table_no]))
            table_no += 1
    complete = (paragraph_no == len(paragraphs) and table_no == len(tables)
                and len(seen_paragraphs) == len(paragraphs)
                and len(seen_tables) == len(tables))
    for position, paragraph in enumerate(paragraphs):
        if position not in seen_paragraphs:
            sequence.append(('paragraph', position, paragraph))
    for position, table in enumerate(tables):
        if position not in seen_tables:
            sequence.append(('table', position, table))
    return sequence, complete


def parse_docx(path, *, document_id, source_version):
    """Extract paragraphs and tables from a DOCX package in document body order."""
    docx = _require('docx', 'python-docx')
    document = docx.Document(str(path))
    blocks, unsupported, index = [], [], 0
    sequence, complete = _docx_body_sequence(document)
    if not complete and 'body_order_not_preserved' not in unsupported:
        unsupported.append('body_order_not_preserved')
    for kind, position, item in sequence:
        if _truncated(index, unsupported):
            break
        if kind == 'paragraph':
            paragraph = item
            text = paragraph.text.strip()
            if not text:
                continue
            style = (getattr(paragraph.style, 'name', '') or '')
            heading = style.lower().startswith('heading')
            level = None
            if heading:
                digits = ''.join(char for char in style if char.isdigit())
                level = int(digits) if digits else 1
            blocks.append(_block(document_id, source_version, 'title' if heading else 'paragraph',
                                 text, {'locator_kind': 'office_object', 'document_id': document_id,
                                        'source_version': source_version, 'part': 'word/document.xml',
                                        'paragraph': position}, index=index, heading_level=level))
            index += 1
        else:
            table = item
            cells = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            text = '\n'.join(' | '.join(row) for row in cells)
            if not text.strip():
                continue
            blocks.append(_block(document_id, source_version, 'table', text,
                                 {'locator_kind': 'office_object', 'document_id': document_id,
                                  'source_version': source_version, 'part': 'word/document.xml',
                                  'table': position}, index=index, table_cells=cells))
            index += 1
    status = 'partial' if unsupported else 'parsed'
    return _result(document_id, DOCX, blocks, adapter='docx-native', parse_status=status,
                   unsupported=unsupported, source_name=Path(path).name)


def parse_pptx(path, *, document_id, source_version):
    """Extract text frames from a PPTX package, slide by slide."""
    pptx = _require('pptx', 'python-pptx')
    presentation = pptx.Presentation(str(path))
    blocks, unsupported, index = [], [], 0
    for slide_no, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        for shape_no, shape in enumerate(slide.shapes):
            if _truncated(index, unsupported):
                break
            if not getattr(shape, 'has_text_frame', False):
                continue
            text = (shape.text_frame.text or '').strip()
            if not text:
                continue
            is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
            blocks.append(_block(document_id, source_version, 'title' if is_title else 'paragraph',
                                 text, {'locator_kind': 'office_object', 'document_id': document_id,
                                        'source_version': source_version,
                                        'part': f'ppt/slides/slide{slide_no}.xml',
                                        'slide': slide_no, 'shape': shape_no},
                                 index=index, heading_level=1 if is_title else None))
            index += 1
    status = 'partial' if unsupported else 'parsed'
    return _result(document_id, PPTX, blocks, adapter='pptx-native', parse_status=status,
                   unsupported=unsupported, source_name=Path(path).name)


def _formula_without_cached_value(formula_row, cached_row):
    """True when a cell holds a formula but no computed value was stored alongside it.

    openpyxl only exposes the formula under ``data_only=False`` and the cached result under
    ``data_only=True``, so both passes are needed to see "a formula lives here" and "but
    nothing was computed for it". A file written by openpyxl (never recalculated by Excel)
    is exactly this shape, and dropping the cell would report an incomplete read as complete.
    """
    for position, cell in enumerate(formula_row):
        value = getattr(cell, 'value', None)
        is_formula = getattr(cell, 'data_type', None) == 'f' or (
            isinstance(value, str) and value.startswith('='))
        if not is_formula:
            continue
        cached = cached_row[position] if position < len(cached_row) else None
        if cached is None:
            return True
    return False


def parse_xlsx(path, *, document_id, source_version, max_rows_per_sheet=500):
    """Extract non-empty sheet rows as ranged blocks; formulas are read as cached values.

    A formula cell without a cached value has no result to cite, so the run is declared
    `partial` with `formula_without_cached_value` rather than reported as a complete read.
    """
    openpyxl = _require('openpyxl', 'openpyxl')
    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    formulas = openpyxl.load_workbook(str(path), data_only=False, read_only=True)
    blocks, unsupported, index = [], [], 0
    try:
        for sheet, formula_sheet in zip(workbook.worksheets, formulas.worksheets):
            value_rows = sheet.iter_rows(values_only=True)
            formula_rows = formula_sheet.iter_rows()
            for row_no, (row, formula_row) in enumerate(
                    zip_longest(value_rows, formula_rows, fillvalue=()), start=1):
                if _truncated(index, unsupported):
                    break
                if row_no > max_rows_per_sheet:
                    if 'sheet_row_limit_reached' not in unsupported:
                        unsupported.append('sheet_row_limit_reached')
                    break
                if _formula_without_cached_value(formula_row, row):
                    if 'formula_without_cached_value' not in unsupported:
                        unsupported.append('formula_without_cached_value')
                values = ['' if value is None else str(value) for value in row]
                if not any(value.strip() for value in values):
                    continue
                last = max(position for position, value in enumerate(values) if value.strip()) + 1
                first_letter = _column_letter(1)
                last_letter = _column_letter(last)
                cell_range = f'{first_letter}{row_no}:{last_letter}{row_no}'
                blocks.append(_block(document_id, source_version, 'table',
                                     ' | '.join(values[:last]).strip(),
                                     {'locator_kind': 'sheet_range', 'document_id': document_id,
                                      'source_version': source_version, 'sheet': sheet.title,
                                      'range': cell_range}, index=index,
                                     table_cells=[values[:last]]))
                index += 1
    finally:
        workbook.close()
        formulas.close()
    status = 'partial' if unsupported else 'parsed'
    return _result(document_id, XLSX, blocks, adapter='xlsx-native', parse_status=status,
                   unsupported=unsupported, source_name=Path(path).name)


def _column_letter(position):
    letters = ''
    while position > 0:
        position, remainder = divmod(position - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def register_office_adapters():
    """Register the three Office adapters with the router (idempotent)."""
    registered = set(parse_contract.registered_adapters())
    handlers = {'docx-native': parse_docx, 'pptx-native': parse_pptx, 'xlsx-native': parse_xlsx}
    capabilities = {'docx-native': ['text', 'headings', 'tables'],
                    'pptx-native': ['text', 'slides'],
                    'xlsx-native': ['text', 'sheet_rows', 'cached_values']}
    for name, media_type in ADAPTERS.items():
        if name in registered:
            continue
        parse_contract.register_adapter(name, media_types=[media_type],
                                        capabilities=capabilities[name], handler=handlers[name])
    return sorted(ADAPTERS)


def parse(path, *, document_id, source_version, **options):
    """Detect the package type and dispatch through the registered Office adapter.

    Extra keyword options are forwarded to the adapter handler (for example
    `max_rows_per_sheet`); an option a handler does not accept surfaces as
    PARSE_RESULT_INVALID rather than being ignored.
    """
    register_office_adapters()
    detected = parse_contract.detect_media_type(path)
    return parse_contract.dispatch(detected, path=path, document_id=document_id,
                                  source_version=source_version, **options)
