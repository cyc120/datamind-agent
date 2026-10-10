# -*- coding: utf-8 -*-
"""WP10 DOCX 构建工具（本地新增；只被交付物生成脚本引用，不改任何既有文件）。

用途：为一页结论卡/设计报告/技术说明书/创新自述提供统一的版式函数。
所有正文数字均来自仓库真实材料（reports/、docs/、只读只读 HTTP/SQL），本模块不产生数字。
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

EA = "PingFang SC"          # macOS 上 LibreOffice 可取到的中文字体
LATIN = "Helvetica Neue"


def _style_run(run, size=9.5, bold=False, mono=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    name = "Menlo" if mono else LATIN
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), EA)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def new_doc(margin_cm=1.8):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(margin_cm)
        section.bottom_margin = Cm(margin_cm)
        section.left_margin = Cm(margin_cm)
        section.right_margin = Cm(margin_cm)
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EA)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.12
    return doc


def para(doc, text="", size=9.5, bold=False, align=None, space_after=3,
         space_before=0, mono=False, color=None, indent=None):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    if text:
        _style_run(p.add_run(text), size=size, bold=bold, mono=mono, color=color)
    return p


def h1(doc, text, size=15):
    p = para(doc, text, size=size, bold=True, space_before=8, space_after=5)
    pPr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F3D7A")
    bdr.append(bottom)
    pPr.append(bdr)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x3D, 0x7A)
    return p


def h2(doc, text, size=11.5):
    p = para(doc, text, size=size, bold=True, space_before=6, space_after=3)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x25, 0x4A, 0x8A)
    return p


def h3(doc, text, size=10):
    p = para(doc, text, size=size, bold=True, space_before=4, space_after=2)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p


def bullets(doc, items, size=9.5, bullet="·", indent=0.35, space_after=2):
    for it in items:
        para(doc, f"{bullet} {it}", size=size, indent=indent, space_after=space_after)


def numbered(doc, items, size=9.5, indent=0.35, space_after=2):
    for i, it in enumerate(items, 1):
        para(doc, f"{i}. {it}", size=size, indent=indent, space_after=space_after)


def code_block(doc, text, size=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.line_spacing = 1.0
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F4F6F8")
    pPr.append(shd)
    for i, line in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        _style_run(p.add_run(line), size=size, mono=True, color=(0x22, 0x33, 0x44))
    return p


def table(doc, header, rows, widths=None, size=8.5, header_fill="DCE6F5"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    hdr = t.rows[0].cells
    for i, text in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        _style_run(p.add_run(str(text)), size=size, bold=True)
        tcPr = hdr[i]._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), header_fill)
        tcPr.append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            _style_run(p.add_run("" if val is None else str(val)), size=size)
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    return t


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def add_footer_pagenum(doc, label):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        _style_run(p.add_run(f"{label}　·　第 "), size=7.5, color=(0x88, 0x88, 0x88))
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "15")
        rPr.append(sz)
        run.append(rPr)
        t = OxmlElement("w:t")
        t.text = "1"
        run.append(t)
        fld.append(run)
        p._element.append(fld)
        _style_run(p.add_run(" 页"), size=7.5, color=(0x88, 0x88, 0x88))


def highlight_box(doc, title, lines, size=9.5, fill="FFF6E5"):
    """单格高亮框（用于亮点预览页）。"""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    _style_run(p.add_run(title), size=size + 0.5, bold=True, color=(0x1F, 0x3D, 0x7A))
    for line in lines:
        pp = cell.add_paragraph()
        pp.paragraph_format.space_after = Pt(0)
        pp.paragraph_format.line_spacing = 1.0
        _style_run(pp.add_run(line), size=size)
    return t
